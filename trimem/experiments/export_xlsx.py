"""Export every TriMEM experiment of this study into one workbook.

    PYTHONPATH=. python -m trimem.experiments.export_xlsx --out triMEM_optimizer/trimem_experiments.xlsx

Sheets (one row per record, raw values, no rounding):
  README                 what each sheet is, settings, status at export time
  fig1b_plans            synthetic Fig. 1(b) query under 5 plans (trimem_dev, decidable predicate)
  evotrace_plan_choice   EvoTrace plan-choice scenarios S1-S3 (planner estimate vs measured)
  greedy_run1_iters      greedy replay, 5 plans x 4 tasks x 10 it (run1; quality invalid, cost kept)
  greedy_run2_iters      greedy replay, 5 plans x 4 tasks x 10 it (run2) + resume to 100 (partial)
  greedy_run3_iters      greedy replay replicate (partial, stopped)
  greedy_offline_iters   offline experience graph, greedy agent, 5 plans x 5 tasks x 2 reps (stopped early)
  truncation             prompt/program/output truncation per cell
  experience_graph_build offline construction cost of the experience graph
  plan_summary           every plan (7 fixed, v1 random, A1-A6, v2 random B0, B1-B3): mean best @10..@50, tokens, validity
  cells_all_plans        one row per task x plan: best @10..@50, costs over iterations 1-20 and 1-50, directories
  curves_0_50            best-so-far per iteration per cell (iterations 0-20, resumed 21-50 stitched)
  children               every generated child (trace + children the OpenEvolve tracer dropped, from checkpoints)
  retrieval_calls        one row per TriMEM retrieval / optimizer decision / observed outcome
  llm_requests           one row per OpenEvolve generation request
  injection_trace        per-iteration injection bookkeeping (requested / rendered / injected / own programs)
  optimizer_decisions    trimem_outcome: every decision (fixed cells backfilled, random, A1-A6, B0-B3) with outcomes
  optimizer_estimators   leave-one-task-out quality of the v1 (delta vs parent) and v2 (delta vs best) priors
  graph_edge_types       native graph edge types (incl. reverse lineage_up / similar_up) and degree statistics
  environment            reproducibility record: hardware, software/model versions, vLLM args, OpenEvolve frozen config,
                         database/experience graph, prompts, math and ALE protocols (also environment.json)
  cell_settings          every cell's resolved settings from its receipt (arm, evaluator, slots, seed, policy, git)
"""

from __future__ import annotations

import argparse
import json
import os
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from trimem.experiments.oe_analyze import cell as oe_cell

R = Path("results/trimem_v0.1.0")


def _flat(d: dict, prefix: str = "") -> dict:
    out = {}
    for k, v in d.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out |= _flat(v, key + ".")
        elif isinstance(v, (list, tuple)):
            out[key] = json.dumps(v) if len(json.dumps(v)) < 30000 else json.dumps(v)[:30000]
        else:
            out[key] = v
    return out


def _sheet(wb: Workbook, name: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(name[:31])
    if not rows:
        ws.append(["(no data)"])
        return
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(c)) for c in cols])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, len(c) + 2))


def _cell(v):
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v


def _jsonl(path: Path) -> list[dict]:
    return [json.loads(x) for x in path.open() if x.strip()] if path.exists() else []


def greedy_iters(run_dir: Path, label: str) -> list[dict]:
    rows = []
    for f in sorted(run_dir.glob("*.jsonl")):
        for r in _jsonl(f):
            flat = _flat({k: v for k, v in r.items() if k not in ("cell",)})
            rows.append({"run": label, "cell": r.get("cell"), **flat})
    return rows


def fig1b() -> list[dict]:
    p = R / "fig1b_plans.json"
    if not p.exists():
        return []
    d = json.loads(p.read_text())
    keep = ("id", "name", "role", "planner_total_cost", "execution_ms", "planning_ms", "warm_wall_s",
            "llm_wait_ms", "rows_judged", "llm_calls", "invalid_replies", "prompt_tokens",
            "completion_tokens", "matches_oracle")
    return [{**{k: x.get(k) for k in keep}, "settings": json.dumps(x.get("settings")),
             "result_ids": json.dumps([r["id"] for r in x["result"]]),
             "oracle_ids": json.dumps([r["id"] for r in x["oracle_result"]]),
             "verdicts_judged": (x.get("verdict_accuracy") or {}).get("judged"),
             "verdicts_correct": (x.get("verdict_accuracy") or {}).get("correct"),
             "model": d.get("model"), "predicate": d.get("predicate")} for x in d["plans"]]


def evotrace_choice() -> list[dict]:
    p = R / "evotrace_plans.json"
    if not p.exists():
        return []
    d = json.loads(p.read_text())
    keep = ("candidate", "planner_total_cost", "execution_ms", "planning_ms", "warm_wall_s", "llm_wait_ms",
            "rows_judged", "worker_requests", "cache_hits", "edges_traversed", "expand_calls", "llm_calls",
            "invalid_replies", "prompt_tokens", "completion_tokens")
    rows = []
    for s in d["scenarios"]:
        for r in [s["baseline"], *s["runs"]]:
            rows.append({"scenario": s["id"], "title": s["title"], "planner_decision": s["auto_decision"],
                         "predicate": s["predicate"], **{k: r.get(k) for k in keep},
                         "top5": json.dumps([x[0] for x in r["result"]])})
    return rows


def openevolve(oe: Path) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    cells, curves, ret, llm = [], [], [], []
    for d in sorted(p for p in oe.iterdir() if p.is_dir() and "__" in p.name):
        task, plan, seed = d.name.split("__")
        receipt = json.loads((d / "run_receipt.json").read_text()) if (d / "run_receipt.json").exists() else {}
        c = oe_cell(d, 20)
        base = {"cell": d.name, "task": f"math:{task}", "plan": plan, "seed": seed}
        if c is None:
            cells.append({**base, "status": "queued"})
            continue
        n = len(c["curve"]) - 1
        cells.append({**base, "status": receipt.get("status"), "iterations_observed": n,
                      "seed_score": c["seed"], "best_at_10": c["curve"][10] if n >= 10 else None,
                      "best_at_20": c["curve"][20] if n >= 20 else None, "best": c["best"],
                      **{k: c[k] for k in ("valid", "generation_ms", "generation_in", "generation_out",
                                           "retrieval_ms", "retrieval_sem_in", "retrieval_sem_out",
                                           "retrieval_rows_judged", "retrieval_worker_requests",
                                           "retrieval_edges", "evaluation_ms")},
                      "placements": json.dumps(c["placements"]),
                      "injected_programs": receipt.get("injected_programs"),
                      "trimem_mode": next((a for a in ("trimodal", "hard", "rank") if plan.startswith(a)), None),
                      "forced_placement": {"semfirst": "d=base", "travfirst": "d=top"}.get(plan.split("-")[-1])})
        for i, v in enumerate(c["curve"]):
            curves.append({**base, "iteration": i, "best_so_far": v})
        for r in _jsonl(d / "retrieval_telemetry.jsonl"):
            ret.append({**base, **_flat(r)})
        for r in _jsonl(d / "llm_telemetry.jsonl"):
            llm.append({**base, **{k: r.get(k) for k in (
                "status", "started_at_unix", "llm_api_ms", "prompt_tokens", "completion_tokens", "total_tokens",
                "finish_reason", "input_chars", "output_chars", "thinking_enabled", "max_tokens_requested")}})
    return cells, curves, ret, llm


def truncation() -> list[dict]:
    paths = sorted((R / "openevolve").glob("truncation_report*.json"))
    if not paths:
        return []
    d = json.loads(paths[-1].read_text())
    return [_flat(r) for r in d.get("greedy", []) + d.get("openevolve", [])]


def build_cost() -> list[dict]:
    rows = []
    for p in sorted((R / "offline").glob("experience_graph_build*.json")):
        d = json.loads(p.read_text())
        for task, v in d.get("per_task", {}).items():
            rows.append({"file": p.name, "task": task, "wall_s_batch": d.get("wall_s"), **v,
                         **{f"census.{k}": x for k, x in d.get("census", {}).get(task, {}).items()}})
    return rows


def outcomes() -> list[dict]:
    import psycopg

    dsn = f"postgresql://{os.environ.get('USER')}@127.0.0.1:55441/trimem_evo"
    with psycopg.connect(dsn) as conn:
        cur = conn.execute("SELECT * FROM trimem_outcome WHERE session NOT LIKE 'smoke%%' ORDER BY session, iteration")
        cols = [d.name for d in cur.description]
        rows = [dict(zip(cols, r, strict=True)) for r in cur.fetchall()]
    for r in rows:
        r["features"] = json.dumps(r["features"])
        r["candidates"] = json.dumps(r["candidates"]) if r["candidates"] is not None else None
        for k in ("decided_at", "observed_at"):
            r[k] = r[k].isoformat() if r[k] else None
    return rows


def estimators() -> list[dict]:
    path = R / "objective" / "priors" / "estimators.json"
    if not path.exists():
        return []
    rep = json.loads(path.read_text())
    return [{"held_out_task": t, **v} for t, v in rep.get("loto", {}).items()]


GROUP = {"nocontext": "fixed", "none": "fixed", "trimodal": "fixed", "rank": "fixed", "hard": "fixed",
         "hard-semfirst": "fixed", "hard-travfirst": "fixed", "adaptive-random": "v1 random exploration",
         "adaptive-v2-random": "v2 random exploration (B0)"}
TAGS = {"adaptive-o1-op": "A1", "adaptive-o1-onc": "A2", "adaptive-o1-onw": "A3", "adaptive-o2-op": "A4",
        "adaptive-o2-onc": "A5", "adaptive-o2-onw": "A6", "adaptive-v2-random": "B0", "adaptive-v2-o3-op": "B1",
        "adaptive-v2-o3-onc": "B2", "adaptive-v2-o3-onw": "B3", "adaptive-random": "R"}
CKS = (10, 20, 30, 40, 50)


def all_plans() -> dict[str, list[dict]]:
    """Every plan's execution data, iterations 1-20 and the resumed 21-50 stitched together."""
    from trimem.experiments.oe_analyze import _checkpoint_children
    from trimem.experiments.render_it20_summary import ADAPTIVE, FIXED, TASKS, construction, row

    cons = construction()
    cells, curves, children, retrieval, llm, injection = [], [], [], [], [], []
    for plan in FIXED + ADAPTIVE:
        base = R / ("objective" if plan.startswith("adaptive") else "openevolve")
        for task in TASKS:
            name = f"{task}__{plan}__s42"
            r = row(base / "it20" / name, task, plan, cons)
            meta = {"task": f"math:{task}", "plan": plan, "tag": TAGS.get(plan, ""),
                    "group": GROUP.get(plan, "v1 optimizer (A1-A6)" if "-o1-" in plan or "-o2-" in plan
                                       else "v2 optimizer (B1-B3)"), "seed": 42}
            if r is None:
                cells.append({**meta, "status": "missing / incomplete"})
                continue
            c50 = r.get("cost50") or {}
            cells.append({**meta, "status": "complete", "iterations": len(r["curve"]) - 1, "seed_score": r["seed"],
                          **{f"best_at_{i}": (r["curve"][i] if i < len(r["curve"]) else None) for i in CKS},
                          "valid_1_20": r["valid"], "gen_in_tokens_1_20": r["gen_in"], "gen_out_tokens_1_20": r["gen_out"],
                          "gen_ms_1_20": r["gen_ms"], "retrieval_ms_1_20": r["ret_ms"], "sem_tokens_1_20": r["sem_tok"],
                          "eval_ms_1_20": r["eval_ms"], "construction_wall_ms_once": r["cons_wall_ms"],
                          "construction_tokens_once": r["cons_tok"],
                          **{f"{k}_1_50": v for k, v in c50.items()},
                          "dir_1_20": str(base / "it20" / name), "dir_21_50": str(base / "it50" / name)})
            for i, v in enumerate(r["curve"]):
                curves.append({**meta, "iteration": i, "best_so_far": v})
            for phase in ("it20", "it50"):
                d = base / phase / name
                if not d.exists():
                    continue
                tr = _jsonl(d / "evolution_trace.jsonl")
                for x in tr:
                    cm, pm = x.get("child_metrics") or {}, x.get("parent_metrics") or {}
                    children.append({**meta, "phase": phase, "iteration": x.get("iteration"), "source": "trace",
                                     "parent_id": x.get("parent_id"), "child_id": x.get("child_id"),
                                     "parent_score": pm.get("combined_score"), "child_score": cm.get("combined_score"),
                                     "valid": (cm.get("combined_score") or 0) > 0, "eval_s": cm.get("eval_time")})
                receipt = json.loads((d / "run_receipt.json").read_text()) if (d / "run_receipt.json").exists() else {}
                extra, _ = _checkpoint_children(d, skip={x.get("child_id") for x in tr})
                for it, sc in extra:
                    if it > int(receipt.get("resumed_at_iteration") or 0):
                        children.append({**meta, "phase": phase, "iteration": it, "source": "checkpoint (dropped by tracer)",
                                         "child_score": sc, "valid": sc > 0})
                for x in _jsonl(d / "retrieval_telemetry.jsonl"):
                    retrieval.append({**meta, "phase": phase, **_flat(x)})
                for x in _jsonl(d / "llm_telemetry.jsonl"):
                    llm.append({**meta, "phase": phase, **{k: x.get(k) for k in (
                        "status", "started_at_unix", "llm_api_ms", "prompt_tokens", "completion_tokens", "total_tokens",
                        "finish_reason", "input_chars", "output_chars", "thinking_enabled", "max_tokens_requested")}})
                for x in _jsonl(d / "injection_trace.jsonl"):
                    injection.append({**meta, "phase": phase, "iteration": x.get("iteration"), "arm": x.get("arm"),
                                      "parent_id": x.get("parent_id"), "requested": x.get("requested"),
                                      "rendered": x.get("rendered"), "budget": x.get("budget"),
                                      "retrieved": x.get("retrieved"), "injected": len(x.get("injected_ids") or []),
                                      "own": len(x.get("own_ids") or []), "error": x.get("error")})
    summary = []
    plans = list(dict.fromkeys(c["plan"] for c in cells))
    for plan in plans:
        cs = [c for c in cells if c["plan"] == plan and c.get("status") == "complete"]
        if not cs:
            continue
        mean = lambda k: (sum(c[k] for c in cs if c.get(k) is not None) / len([c for c in cs if c.get(k) is not None])  # noqa: E731
                          if any(c.get(k) is not None for c in cs) else None)
        summary.append({"plan": plan, "tag": cs[0]["tag"], "group": cs[0]["group"], "tasks": len(cs),
                        **{f"mean_best_at_{i}": mean(f"best_at_{i}") for i in CKS},
                        "mean_valid_rate_1_20": mean("valid_1_20") / 20,
                        "mean_tokens_1_20": sum(c["gen_in_tokens_1_20"] + c["gen_out_tokens_1_20"] + c["sem_tokens_1_20"] for c in cs) / len(cs),
                        "mean_tokens_1_50": (sum(c["gen_in_1_50"] + c["gen_out_1_50"] + c["sem_tok_1_50"] for c in cs) / len(cs))
                        if all("gen_in_1_50" in c for c in cs) else None,
                        "mean_valid_rate_1_50": (sum(c["valid_1_50"] for c in cs) / len(cs) / 50) if all("valid_1_50" in c for c in cs) else None,
                        "mean_gen_ms_1_20": mean("gen_ms_1_20"), "mean_retrieval_ms_1_20": mean("retrieval_ms_1_20")})
    return {"plan_summary": summary, "cells_all_plans": cells, "curves_0_50": curves, "children": children,
            "retrieval_calls": retrieval, "llm_requests": llm, "injection_trace": injection}


def estimators_all() -> list[dict]:
    rows = []
    for name in ("priors", "priors_v2"):
        path = R / "objective" / name / "estimators.json"
        if path.exists():
            rep = json.loads(path.read_text())
            for t, v in rep.get("loto", {}).items():
                rows.append({"prior": name, "label": rep.get("label", "delta (vs parent)"), "n_train_total": rep.get("n"),
                             "held_out_task": t, **v})
    return rows


def graph_stats() -> list[dict]:
    import psycopg

    dsn = f"postgresql://{os.environ.get('USER')}@127.0.0.1:55441/trimem_evo"
    with psycopg.connect(dsn) as conn:
        cur = conn.execute("SELECT * FROM trimem.graph_stats ORDER BY type_id")
        cols = [d.name for d in cur.description]
        return [{k: (str(v) if k == "analyzed_at" else v) for k, v in zip(cols, r, strict=True)} for r in cur.fetchall()]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="triMEM_optimizer/trimem_experiments.xlsx")
    args = ap.parse_args(argv)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    oe_cells, oe_curves, oe_ret, oe_llm = openevolve(R / "openevolve" / "it20")
    status = {}
    for c in oe_cells:
        status[c.get("status")] = status.get(c.get("status"), 0) + 1
    lines = [
        ["TriMEM experiments workbook", time.strftime("%Y-%m-%d %H:%M:%S")],
        ["Generator", "Qwen/Qwen3.8-27B-FP8 (vLLM, thinking off); embeddings Qwen3-Embedding-0.6B"],
        ["Memory", "EvoTrace experience graph in PostgreSQL 16 (trimem_evo): attempts, HNSW, native lineage + "
                   "similarity graphs, LLM change summaries"],
        ["Semantic predicate", "concrete transferable algorithmic idea (hard filter) / relevance score 0-10 (rank)"],
        ["OpenEvolve status at export", json.dumps(status)],
        ["Note", "Cells still running are read from their latest OpenEvolve checkpoint; best_at_20 is empty "
                 "until a cell reaches iteration 20. Greedy offline cells were stopped early by request."],
        [],
        ["Sheet", "Contents"],
    ]
    for line in (__doc__ or "").split("Sheets (one row per record, raw values, no rounding):")[1].strip().splitlines():
        name, _, desc = line.strip().partition(" ")
        lines.append([name, desc.strip()])
    for line in lines:
        readme.append(line)
    readme.column_dimensions["A"].width = 26
    readme.column_dimensions["B"].width = 120
    _sheet(wb, "fig1b_plans", fig1b())
    _sheet(wb, "evotrace_plan_choice", evotrace_choice())
    for run in ("run1", "run2", "run3"):
        _sheet(wb, f"greedy_{run}_iters", greedy_iters(R / "replay" / run, run))
    offline = []
    for rep in ("greedy_it20_rep1", "greedy_it20_rep2"):
        offline += greedy_iters(R / "offline" / rep, rep)
    _sheet(wb, "greedy_offline_iters", offline)
    _sheet(wb, "truncation", truncation())
    _sheet(wb, "experience_graph_build", build_cost())
    for name, rows in all_plans().items():
        _sheet(wb, name, rows)
    _sheet(wb, "optimizer_decisions", outcomes())
    _sheet(wb, "optimizer_estimators", estimators_all())
    _sheet(wb, "graph_edge_types", graph_stats())
    from trimem.experiments.environment import cells as env_cells
    from trimem.experiments.environment import collect as env_collect

    env = env_collect()
    Path(out).with_name("environment.json").write_text(json.dumps(env, indent=1, default=str))
    env_rows = [{"section": k.split(".", 1)[0], "key": k, "value": v if not isinstance(v, (list, dict)) else json.dumps(v)}
                for k, v in _flat(env).items()]
    _sheet(wb, "environment", env_rows)
    _sheet(wb, "cell_settings", env_cells())
    wb.save(out)
    print(out, {ws.title: ws.max_row - 1 for ws in wb.worksheets})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
