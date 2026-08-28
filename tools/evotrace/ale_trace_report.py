"""Offline ALE trace diagnostics and workbook augmentation.

This module derives the trace-only analyses used by arXiv:2605.20086 without
rerunning evolutionary search:

* best-so-far LOC and numeric-literal trajectories;
* final-best discovery time and lineage depth;
* deterministic code-line recycling (cycling);
* a transparent rule-based proxy for the paper's LLM edit taxonomy; and
* descriptive comparisons between no-memory, GEM, and Polyglot configurations.

The taxonomy proxy is intentionally named as such everywhere.  It is useful for
screening the local corpus, but is not interchangeable with the paper's validated
LLM-as-judge labels.  Every plotted value is also exported as CSV and embedded in
the requested workbook for auditability.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import json
import math
import os
import re
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

SCHEMA_VERSION = "tridb_ale_trace_report_v0.1.0"
PAPER = "Pelleriti et al., arXiv:2605.20086, sections 5.1-5.2 and B.2-B.6"
EXPECTED_TASKS = {
    "ahc008",
    "ahc011",
    "ahc015",
    "ahc016",
    "ahc024",
    "ahc025",
    "ahc026",
    "ahc027",
    "ahc039",
    "ahc046",
}
CONFIG_ORDER = [
    "No Memory",
    "GEM p=0.1",
    "GEM p=0.5",
    "GEM p=1",
    "Polyglot p=0.1",
    "Polyglot p=0.5",
    "Polyglot p=1",
]
BASE_ORDER = ["No Memory", "GEM", "Polyglot"]
COLORS = {"No Memory": "#777777", "GEM": "#2878B5", "Polyglot": "#E07A1F"}
LABELS = [
    "hyperparameter_tuning",
    "local_refinement",
    "architectural_change",
    "composition",
    "efficiency",
    "bug_fix",
    "external_dependency",
    "pruning",
    "refactor",
]
LABEL_DISPLAY = {
    "hyperparameter_tuning": "Hyperparameter tuning",
    "local_refinement": "Local refinement",
    "architectural_change": "Architectural change",
    "composition": "Composition",
    "efficiency": "Efficiency",
    "bug_fix": "Bug fix",
    "external_dependency": "External dependency",
    "pruning": "Pruning",
    "refactor": "Refactor",
}
SURFACE = "#fcfcfb"
GRID = "#d9dee7"
INK = "#111111"
MUTED = "#555555"

NUMBER_RE = re.compile(
    r"(?<![A-Za-z0-9_\.])(?:0[xX][0-9A-Fa-f]+|0[bB][01]+|"
    r"(?:\d+\.\d*|\.\d+|\d+)(?:[eE][+-]?\d+)?)(?:[uUlLfF]+)?"
)
STRING_RE = re.compile(r'"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'')
BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
LINE_COMMENT_RE = re.compile(r"//.*?$", re.MULTILINE)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"expected JSON object: {path}")
    return value


def _config(receipt: Mapping[str, Any]) -> tuple[str, str]:
    arm = str(receipt["arm"])
    if arm == "nocontext":
        return "No Memory", "No Memory"
    base = "GEM" if arm == "gem" else "Polyglot" if arm == "polyglot" else arm
    frequency = float(receipt.get("injection_frequency") or 1.0)
    return base, f"{base} p={frequency:g}"


def _strip_for_metrics(code: str) -> str:
    code = BLOCK_COMMENT_RE.sub("", code)
    code = LINE_COMMENT_RE.sub("", code)
    return code


def code_shape(code: str) -> tuple[int, int]:
    """Return non-comment LOC and numeric literal occurrences for C++ source."""
    stripped = _strip_for_metrics(code)
    loc = sum(bool(line.strip()) for line in stripped.splitlines())
    without_strings = STRING_RE.sub('""', stripped)
    return loc, len(NUMBER_RE.findall(without_strings))


def _number_skeleton(line: str) -> str:
    return NUMBER_RE.sub("<NUM>", line)


def _semantic_skeleton(code: str) -> str:
    code = _strip_for_metrics(code)
    code = STRING_RE.sub('"STR"', code)
    code = NUMBER_RE.sub("<NUM>", code)
    return re.sub(r"\s+", "", code)


def _code_lines(code: str) -> list[str]:
    return code.splitlines(keepends=True)


def line_diff(parent: str, child: str) -> tuple[list[str], list[str]]:
    """Return added and deleted physical lines, retaining whitespace and newline."""
    before, after = _code_lines(parent), _code_lines(child)
    matcher = difflib.SequenceMatcher(a=before, b=after, autojunk=False)
    added: list[str] = []
    deleted: list[str] = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag in {"delete", "replace"}:
            deleted.extend(before[i1:i2])
        if tag in {"insert", "replace"}:
            added.extend(after[j1:j2])
    return added, deleted


def _nonblank(lines: Iterable[str]) -> list[str]:
    return [line for line in lines if line.strip()]


def _comment_only(line: str) -> bool:
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith("//")
        or stripped.startswith("/*")
        or stripped.startswith("*")
        or stripped.endswith("*/")
    )


def _programs(cell_dir: Path) -> dict[str, dict[str, Any]]:
    programs: dict[str, dict[str, Any]] = {}
    for path in sorted(
        (cell_dir / "openevolve" / "checkpoints").glob("checkpoint_*/programs/*.json")
    ):
        obj = _read_json(path)
        program_id = obj.get("id")
        if program_id:
            programs[str(program_id)] = obj
    if not programs:
        raise ValueError(f"no checkpoint programs: {cell_dir}")
    return programs


def _trace(cell_dir: Path) -> list[dict[str, Any]]:
    path = cell_dir / "evolution_trace.jsonl"
    return sorted(
        [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ],
        key=lambda row: (int(row["iteration"]), str(row["child_id"])),
    )


def _root(
    programs: Mapping[str, Mapping[str, Any]],
    initial: str,
    trace: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    candidates = [
        dict(program)
        for program in programs.values()
        if not program.get("parent_id")
        and int(program.get("iteration_found") or 0) == 0
    ]
    matching = [
        p for p in candidates if str(p.get("code", "")).strip() == initial.strip()
    ]
    if not matching:
        raise ValueError("no checkpoint seed matches initial_program.py")
    # The memory runs may carry several root variants with byte-identical seed
    # code.  The causal root is the one actually selected as a trace parent.
    trace_parents = {str(row["parent_id"]) for row in trace}
    selected = [program for program in matching if str(program["id"]) in trace_parents]
    if selected:
        selected.sort(key=lambda program: str(program["id"]))
        return selected[0]
    matching.sort(key=lambda program: str(program["id"]))
    return matching[0]


def _lineage(
    best_id: str, programs: Mapping[str, Mapping[str, Any]]
) -> tuple[list[str], bool, str | None]:
    ids: list[str] = []
    seen: set[str] = set()
    current: str | None = best_id
    while current:
        if current in seen:
            return ids, False, f"cycle:{current}"
        seen.add(current)
        program = programs.get(current)
        if program is None:
            return ids, False, f"missing:{current}"
        ids.append(current)
        parent = program.get("parent_id")
        current = None if not parent else str(parent)
    return ids, True, None


def _best_iteration(
    best: Mapping[str, Any],
    trace: Sequence[Mapping[str, Any]],
    programs: Mapping[str, Mapping[str, Any]],
) -> tuple[int, str]:
    best_id = str(best["id"])
    for row in trace:
        if str(row["child_id"]) == best_id:
            return int(row["iteration"]), "trace_child"
    program = programs.get(best_id, best)
    found = program.get("iteration_found")
    if found is not None:
        return int(found), "checkpoint_iteration_found"
    if not program.get("parent_id"):
        return 0, "seed"
    raise ValueError(f"cannot recover final-best iteration for {best_id}")


def load_cells(run_dir: Path) -> list[dict[str, Any]]:
    cells: list[dict[str, Any]] = []
    for cell_dir in sorted(run_dir.glob("ale_*")):
        if not cell_dir.is_dir() or ".outage_" in cell_dir.name:
            continue
        receipt_path = cell_dir / "run_receipt.json"
        if not receipt_path.exists():
            continue
        receipt = _read_json(receipt_path)
        if receipt.get("status") != "complete":
            continue
        trace = _trace(cell_dir)
        programs = _programs(cell_dir)
        initial = (cell_dir / "initial_program.py").read_text(encoding="utf-8")
        root = _root(programs, initial, trace)
        best = _read_json(cell_dir / "openevolve" / "best" / "best_program_info.json")
        lineage_ids, lineage_complete, lineage_error = _lineage(
            str(best["id"]), programs
        )
        best_iteration, best_iteration_source = _best_iteration(best, trace, programs)
        base_arm, config = _config(receipt)
        cells.append(
            {
                "cell": cell_dir.name,
                "dir": cell_dir,
                "task": str(receipt["task_uid"]).split(":", 1)[-1],
                "arm": str(receipt["arm"]),
                "base_arm": base_arm,
                "config": config,
                "frequency": None
                if receipt["arm"] == "nocontext"
                else float(receipt.get("injection_frequency") or 1.0),
                "receipt": receipt,
                "trace": trace,
                "programs": programs,
                "root": root,
                "best": best,
                "lineage_ids": lineage_ids,
                "lineage_complete": lineage_complete,
                "lineage_error": lineage_error,
                "best_iteration": best_iteration,
                "best_iteration_source": best_iteration_source,
            }
        )
    if len(cells) != 70:
        raise ValueError(f"expected 70 complete canonical cells, found {len(cells)}")
    if {cell["task"] for cell in cells} != EXPECTED_TASKS:
        raise ValueError("task coverage does not match the 10 ALE tasks")
    counts = Counter(cell["config"] for cell in cells)
    if any(counts[name] != 10 for name in CONFIG_ORDER):
        raise ValueError(f"configuration coverage is not 10 cells each: {counts}")
    return cells


def shape_trajectories(
    cells: Sequence[dict[str, Any]],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    long_rows: list[dict[str, Any]] = []
    cell_rows: list[dict[str, Any]] = []
    for cell in cells:
        root = cell["root"]
        seed_code = str(root["code"])
        seed_loc, seed_literals = code_shape(seed_code)
        root_score = float((root.get("metrics") or {}).get("combined_score") or 0.0)
        events: dict[int, list[tuple[float, str, str]]] = defaultdict(list)
        for row in cell["trace"]:
            score = (row.get("child_metrics") or {}).get("combined_score")
            if score is None:
                continue
            events[int(row["iteration"])].append(
                (float(score), str(row["child_code"]), str(row["child_id"]))
            )
        # Three final winners were checkpoint-only.  Inserting the saved winner at
        # iteration_found makes the final point auditable without inventing missing edits.
        best_id = str(cell["best"]["id"])
        if not any(
            best_id == child_id
            for values in events.values()
            for _, _, child_id in values
        ):
            best_program = cell["programs"][best_id]
            best_score = float(
                (best_program.get("metrics") or {}).get("combined_score") or 0.0
            )
            events[cell["best_iteration"]].append(
                (best_score, str(best_program["code"]), best_id)
            )

        budget = int(cell["receipt"].get("iterations_expected") or 40)
        best_score, best_code, best_id_running = root_score, seed_code, str(root["id"])
        for iteration in range(budget + 1):
            for score, code, child_id in events.get(iteration, []):
                if score > best_score:
                    best_score, best_code, best_id_running = score, code, child_id
            loc, literals = code_shape(best_code)
            long_rows.append(
                {
                    "cell": cell["cell"],
                    "task": cell["task"],
                    "base_arm": cell["base_arm"],
                    "config": cell["config"],
                    "iteration": iteration,
                    "normalized_iteration": iteration / budget,
                    "best_program_id": best_id_running,
                    "best_public_score": best_score,
                    "loc": loc,
                    "numeric_literals": literals,
                    "loc_ratio_to_seed": loc / seed_loc if seed_loc else None,
                    "numeric_literal_ratio_to_seed": literals / seed_literals
                    if seed_literals
                    else None,
                }
            )
        final_code = str(cell["programs"][best_id]["code"])
        final_loc, final_literals = code_shape(final_code)
        cell_rows.append(
            {
                "cell": cell["cell"],
                "task": cell["task"],
                "base_arm": cell["base_arm"],
                "config": cell["config"],
                "trace_rows": len(cell["trace"]),
                "iterations_completed": int(
                    cell["receipt"].get("iterations_traced") or 0
                ),
                "trace_row_coverage": len(cell["trace"])
                / int(cell["receipt"].get("iterations_traced") or 40),
                "seed_loc": seed_loc,
                "final_loc": final_loc,
                "final_loc_ratio": final_loc / seed_loc if seed_loc else None,
                "seed_numeric_literals": seed_literals,
                "final_numeric_literals": final_literals,
                "final_numeric_literal_ratio": final_literals / seed_literals
                if seed_literals
                else None,
                "best_iteration": cell["best_iteration"],
                "best_normalized_iteration": cell["best_iteration"]
                / int(cell["receipt"].get("iterations_expected") or 40),
                "budget_remaining_fraction": 1
                - cell["best_iteration"]
                / int(cell["receipt"].get("iterations_expected") or 40),
                "best_iteration_source": cell["best_iteration_source"],
                "lineage_depth": len(cell["lineage_ids"]) - 1
                if cell["lineage_complete"]
                else None,
                "known_lineage_depth": len(cell["lineage_ids"]) - 1,
                "lineage_complete": cell["lineage_complete"],
                "lineage_error": cell["lineage_error"],
            }
        )
    return pd.DataFrame(long_rows), pd.DataFrame(cell_rows)


def _history_deleted(
    program_id: str,
    programs: Mapping[str, Mapping[str, Any]],
    cache: dict[str, tuple[set[str], set[str], set[str], bool]],
    active: set[str] | None = None,
) -> tuple[set[str], set[str], set[str], bool]:
    """Deleted-line history before ``program_id`` along its causal lineage."""
    if program_id in cache:
        exact, numeric, trivial, complete = cache[program_id]
        return set(exact), set(numeric), set(trivial), complete
    active = set() if active is None else set(active)
    if program_id in active:
        return set(), set(), set(), False
    active.add(program_id)
    program = programs.get(program_id)
    if program is None:
        return set(), set(), set(), False
    parent_id = program.get("parent_id")
    if not parent_id:
        result = (set(), set(), set(), True)
        cache[program_id] = result
        return set(), set(), set(), True
    parent = programs.get(str(parent_id))
    if parent is None:
        return set(), set(), set(), False
    exact, numeric, trivial, complete = _history_deleted(
        str(parent_id), programs, cache, active
    )
    _, deleted = line_diff(str(parent.get("code", "")), str(program.get("code", "")))
    for line in _nonblank(deleted):
        exact.add(line)
        numeric.add(_number_skeleton(line))
        if _comment_only(line):
            trivial.add(re.sub(r"\s+", "", line))
    cache[program_id] = (set(exact), set(numeric), set(trivial), complete)
    return exact, numeric, trivial, complete


def taxonomy_proxy(
    parent: str,
    child: str,
    added: Sequence[str],
    deleted: Sequence[str],
    description: str,
) -> set[str]:
    """Transparent multi-label screening proxy; not the paper's LLM judge."""
    labels: set[str] = set()
    desc = description.lower()
    add_nonblank, del_nonblank = _nonblank(added), _nonblank(deleted)
    changed = len(add_nonblank) + len(del_nonblank)
    parent_loc = max(code_shape(parent)[0], 1)
    edit_ratio = changed / parent_loc
    parent_numbers = NUMBER_RE.findall(STRING_RE.sub('""', _strip_for_metrics(parent)))
    child_numbers = NUMBER_RE.findall(STRING_RE.sub('""', _strip_for_metrics(child)))

    if parent_numbers != child_numbers and (
        _semantic_skeleton(parent) == _semantic_skeleton(child)
        or re.search(
            r"parameter|constant|threshold|penalty|weight|bonus|depth|limit|tunable",
            desc,
        )
    ):
        labels.add("hyperparameter_tuning")
    if any(re.match(r"\s*#\s*include\b", line) for line in add_nonblank) or re.search(
        r"dependency|library|include|import", desc
    ):
        labels.add("external_dependency")
    if re.search(
        r"efficien|faster|performance|cache|precomput|complexity|avoid.*loop|"
        r"unordered_|priority_queue|bitset|incremental|optimi[sz]e",
        desc + "\n" + "".join(add_nonblank),
    ):
        labels.add("efficiency")
    if re.search(
        r"bug|fix|correct|invalid|bounds?|overflow|runtime|crash|edge case|"
        r"off.by.one|uninitiali[sz]ed|safety",
        desc,
    ):
        labels.add("bug_fix")
    function_adds = sum(
        bool(
            re.search(
                r"\b(?:void|bool|int|long|double|float|char|auto|Point|vector<[^>]+>)\s+"
                r"[A-Za-z_]\w*\s*\([^;]*\)\s*\{",
                line,
            )
        )
        for line in add_nonblank
    )
    if function_adds or re.search(
        r"compos|combine|integrat|helper function|new function", desc
    ):
        labels.add("composition")
    if edit_ratio >= 0.18 or re.search(
        r"architect|redesign|rewrite|new strategy|replace.*algorithm|major overhaul|"
        r"state machine|new (?:class|struct|enum)",
        desc,
    ):
        labels.add("architectural_change")
    if (
        len(del_nonblank) >= 5 and len(del_nonblank) > 1.5 * max(len(add_nonblank), 1)
    ) or re.search(
        r"prun|remove|delete|drop|eliminate|strip out|simplif",
        desc,
    ):
        labels.add("pruning")
    if re.search(
        r"refactor|rename|extract|reorgani[sz]e|cleanup|clean up|helper", desc
    ):
        labels.add("refactor")
    if re.search(r"refin|adjust|improv|local|minor|tweak", desc) or not labels:
        labels.add("local_refinement")
    return labels


def edit_diagnostics(
    cells: Sequence[dict[str, Any]], cell_summary: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    edit_rows: list[dict[str, Any]] = []
    cycling_rows: list[dict[str, Any]] = []
    for cell in cells:
        programs = cell["programs"]
        history_cache: dict[str, tuple[set[str], set[str], set[str], bool]] = {}
        final_lineage = set(cell["lineage_ids"]) if cell["lineage_complete"] else set()
        root_score = float(
            (cell["root"].get("metrics") or {}).get("combined_score") or 0.0
        )
        running_best = root_score
        cumulative_added = cumulative_exact = cumulative_tuning = cumulative_trivial = 0
        for row in cell["trace"]:
            parent_id, child_id = str(row["parent_id"]), str(row["child_id"])
            parent, child = str(row["parent_code"]), str(row["child_code"])
            added, deleted = line_diff(parent, child)
            added_nb, deleted_nb = _nonblank(added), _nonblank(deleted)
            hist_exact, hist_numeric, hist_trivial, history_complete = _history_deleted(
                parent_id, programs, history_cache
            )
            literal = tuning = trivial = 0
            for line in added_nb:
                if line in hist_exact:
                    literal += 1
                elif NUMBER_RE.search(line) and _number_skeleton(line) in hist_numeric:
                    tuning += 1
                elif _comment_only(line) and re.sub(r"\s+", "", line) in hist_trivial:
                    trivial += 1
            recycled = literal + tuning + trivial
            cumulative_added += len(added_nb)
            cumulative_exact += literal
            cumulative_tuning += tuning
            cumulative_trivial += trivial
            parent_score = (row.get("parent_metrics") or {}).get("combined_score")
            child_score = (row.get("child_metrics") or {}).get("combined_score")
            parent_score_f = None if parent_score is None else float(parent_score)
            child_score_f = None if child_score is None else float(child_score)
            positive = bool(
                parent_score_f is not None
                and child_score_f is not None
                and child_score_f > parent_score_f
            )
            best_update = bool(
                child_score_f is not None and child_score_f > running_best
            )
            if child_score_f is not None:
                running_best = max(running_best, child_score_f)
            description = (
                str(row.get("parent_changes_description") or "")
                + "\n"
                + str((row.get("metadata") or {}).get("changes") or "")
            )
            labels = taxonomy_proxy(parent, child, added_nb, deleted_nb, description)
            base = {
                "cell": cell["cell"],
                "task": cell["task"],
                "base_arm": cell["base_arm"],
                "config": cell["config"],
                "iteration": int(row["iteration"]),
                "normalized_iteration": int(row["iteration"])
                / int(cell["receipt"].get("iterations_expected") or 40),
                "parent_id": parent_id,
                "child_id": child_id,
                "parent_score": parent_score_f,
                "child_score": child_score_f,
                "score_delta": None
                if parent_score_f is None or child_score_f is None
                else child_score_f - parent_score_f,
                "positive_score_change": positive,
                "best_so_far_update": best_update,
                "on_complete_final_lineage": child_id in final_lineage,
                "added_nonblank_lines": len(added_nb),
                "deleted_nonblank_lines": len(deleted_nb),
                "labels": ";".join(label for label in LABELS if label in labels),
                "label_count": len(labels),
            }
            for label in LABELS:
                base[label] = label in labels
            edit_rows.append(base)
            cycling_rows.append(
                {
                    "cell": cell["cell"],
                    "task": cell["task"],
                    "base_arm": cell["base_arm"],
                    "config": cell["config"],
                    "iteration": int(row["iteration"]),
                    "normalized_iteration": int(row["iteration"])
                    / int(cell["receipt"].get("iterations_expected") or 40),
                    "parent_id": parent_id,
                    "child_id": child_id,
                    "history_complete": history_complete,
                    "added_nonblank_lines": len(added_nb),
                    "literal_recycled_lines": literal,
                    "tuning_recycled_lines": tuning,
                    "trivial_recycled_lines": trivial,
                    "recycled_lines": recycled,
                    "edge_recycling_share": recycled / len(added_nb)
                    if added_nb
                    else None,
                    "cumulative_added_lines": cumulative_added,
                    "cumulative_recycled_lines": cumulative_exact
                    + cumulative_tuning
                    + cumulative_trivial,
                    "cumulative_recycling_share": (
                        cumulative_exact + cumulative_tuning + cumulative_trivial
                    )
                    / cumulative_added
                    if cumulative_added
                    else None,
                }
            )

        sub = [row for row in cycling_rows if row["cell"] == cell["cell"]]
        added = sum(
            row["added_nonblank_lines"] for row in sub if row["history_complete"]
        )
        exact = sum(
            row["literal_recycled_lines"] for row in sub if row["history_complete"]
        )
        tuning = sum(
            row["tuning_recycled_lines"] for row in sub if row["history_complete"]
        )
        trivial = sum(
            row["trivial_recycled_lines"] for row in sub if row["history_complete"]
        )
        points = [
            (row["normalized_iteration"], row["cumulative_recycling_share"])
            for row in sub
            if row["history_complete"] and row["cumulative_recycling_share"] is not None
        ]
        slope = None
        if len(points) >= 3:
            slope = float(
                np.polyfit([x for x, _ in points], [y for _, y in points], 1)[0]
            )
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "cycling_history_complete_fraction"
        ] = sum(row["history_complete"] for row in sub) / len(sub) if sub else None
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "added_lines_with_history"
        ] = added
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "literal_recycling_share"
        ] = exact / added if added else None
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "tuning_recycling_share"
        ] = tuning / added if added else None
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "trivial_recycling_share"
        ] = trivial / added if added else None
        cell_summary.loc[cell_summary["cell"] == cell["cell"], "cycling_share"] = (
            (exact + tuning + trivial) / added if added else None
        )
        cell_summary.loc[cell_summary["cell"] == cell["cell"], "cycling_slope"] = slope
        edit_sub = [row for row in edit_rows if row["cell"] == cell["cell"]]
        cell_summary.loc[
            cell_summary["cell"] == cell["cell"], "positive_edit_fraction"
        ] = (
            sum(row["positive_score_change"] for row in edit_sub) / len(edit_sub)
            if edit_sub
            else None
        )
        cell_summary.loc[cell_summary["cell"] == cell["cell"], "best_update_count"] = (
            sum(row["best_so_far_update"] for row in edit_sub)
        )
    return pd.DataFrame(cycling_rows), pd.DataFrame(edit_rows)


def odds_ratio(
    values: Sequence[bool], outcomes: Sequence[bool]
) -> tuple[float, float, float]:
    a = sum(v and o for v, o in zip(values, outcomes, strict=True))
    b = sum(v and not o for v, o in zip(values, outcomes, strict=True))
    c = sum(not v and o for v, o in zip(values, outcomes, strict=True))
    d = sum(not v and not o for v, o in zip(values, outcomes, strict=True))
    a, b, c, d = (x + 0.5 for x in (a, b, c, d))
    ratio = a * d / (b * c)
    se = math.sqrt(1 / a + 1 / b + 1 / c + 1 / d)
    return (
        ratio,
        math.exp(math.log(ratio) - 1.96 * se),
        math.exp(math.log(ratio) + 1.96 * se),
    )


def taxonomy_summary(edits: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    outcomes = edits["positive_score_change"].astype(bool).tolist()
    for label in LABELS:
        present = edits[label].astype(bool).tolist()
        ratio, low, high = odds_ratio(present, outcomes)
        all_rate = float(edits[label].mean())
        best = edits[edits["best_so_far_update"]]
        lineage = edits[edits["on_complete_final_lineage"]]
        rows.append(
            {
                "label": label,
                "display_label": LABEL_DISPLAY[label],
                "labeled_edits": int(edits[label].sum()),
                "prevalence": all_rate,
                "positive_change_odds_ratio": ratio,
                "odds_ratio_ci_low": low,
                "odds_ratio_ci_high": high,
                "best_update_enrichment": float(best[label].mean() / all_rate)
                if len(best) and all_rate
                else None,
                "final_lineage_enrichment": float(lineage[label].mean() / all_rate)
                if len(lineage) and all_rate
                else None,
                **{
                    f"prevalence_{base.lower().replace(' ', '_')}": float(
                        edits[edits["base_arm"] == base][label].mean()
                    )
                    for base in BASE_ORDER
                },
            }
        )
    return pd.DataFrame(rows)


def arm_summary(cell_summary: pd.DataFrame, edits: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for config in CONFIG_ORDER:
        sub = cell_summary[cell_summary["config"] == config]
        esub = edits[edits["config"] == config]
        rows.append(
            {
                "config": config,
                "cells": len(sub),
                "trace_rows": int(sub["trace_rows"].sum()),
                "median_trace_row_coverage": float(sub["trace_row_coverage"].median()),
                "median_final_loc_ratio": float(sub["final_loc_ratio"].median()),
                "median_final_numeric_literal_ratio": float(
                    sub["final_numeric_literal_ratio"].median()
                ),
                "median_best_normalized_iteration": float(
                    sub["best_normalized_iteration"].median()
                ),
                "median_budget_remaining_fraction": float(
                    sub["budget_remaining_fraction"].median()
                ),
                "complete_lineages": int(sub["lineage_complete"].sum()),
                "median_lineage_depth_complete": float(sub["lineage_depth"].median()),
                "median_cycling_share": float(sub["cycling_share"].median()),
                "positive_cycling_slope_cells": int((sub["cycling_slope"] > 0).sum()),
                "mean_positive_edit_fraction": float(
                    esub["positive_score_change"].mean()
                ),
                "best_so_far_updates": int(esub["best_so_far_update"].sum()),
            }
        )
    return pd.DataFrame(rows)


def _style_axis(ax: Any, grid_axis: str = "both") -> None:
    ax.set_facecolor(SURFACE)
    ax.grid(axis=grid_axis, color=GRID, linewidth=0.8, alpha=0.8)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(colors=MUTED, labelsize=8)


def _save(fig: Any, output_dir: Path, stem: str) -> list[Path]:
    paths = []
    for suffix, kwargs in (("png", {"dpi": 240}), ("pdf", {})):
        path = output_dir / f"{stem}.{suffix}"
        fig.savefig(path, bbox_inches="tight", facecolor=SURFACE, **kwargs)
        paths.append(path)
    plt.close(fig)
    return paths


def figure_shape(shape: pd.DataFrame, output_dir: Path) -> list[Path]:
    fig, axes = plt.subplots(1, 2, figsize=(10.8, 4.2), constrained_layout=True)
    for ax, metric, title in (
        (axes[0], "loc_ratio_to_seed", "Best-so-far LOC / seed LOC"),
        (
            axes[1],
            "numeric_literal_ratio_to_seed",
            "Best-so-far numeric literals / seed",
        ),
    ):
        for base in BASE_ORDER:
            sub = shape[shape["base_arm"] == base]
            pivot = sub.pivot(index="cell", columns="iteration", values=metric)
            median = pivot.median(axis=0)
            q1, q3 = pivot.quantile(0.25, axis=0), pivot.quantile(0.75, axis=0)
            x = median.index.to_numpy(dtype=float) / 40.0
            ax.plot(
                x, median, label=f"{base} (n={len(pivot)})", color=COLORS[base], lw=2
            )
            ax.fill_between(
                x,
                q1.to_numpy(float),
                q3.to_numpy(float),
                color=COLORS[base],
                alpha=0.16,
            )
        ax.axhline(1.0, color="#999999", ls="--", lw=1)
        ax.set_xlabel("Normalized search iteration")
        ax.set_ylabel("Ratio to seed")
        ax.set_title(title)
        _style_axis(ax)
    axes[0].legend(frameon=False, fontsize=8)
    fig.suptitle(
        "Figure 13: ALE program shape across the existing 40-iteration traces\n"
        "solid = cross-cell median; band = IQR; missing trace rows are forward-filled",
        fontsize=11,
        weight="semibold",
    )
    return _save(fig, output_dir, "figure13_program_shape")


def _box_strip(
    ax: Any, groups: list[np.ndarray], labels: list[str], ylabel: str
) -> None:
    ax.boxplot(
        groups,
        tick_labels=labels,
        showfliers=False,
        widths=0.58,
        medianprops={"color": "#111111", "linewidth": 1.4},
    )
    rng = np.random.default_rng(42)
    for i, values in enumerate(groups, start=1):
        x = rng.normal(i, 0.055, len(values))
        ax.scatter(x, values, s=16, color="#3e6a8d", alpha=0.62, zorder=3)
    ax.set_ylabel(ylabel)
    ax.tick_params(axis="x", rotation=35)
    _style_axis(ax, "y")


def figure_budget_lineage(cell_summary: pd.DataFrame, output_dir: Path) -> list[Path]:
    fig, axes = plt.subplots(1, 2, figsize=(11.4, 4.6), constrained_layout=True)
    timing = [
        (
            cell_summary[cell_summary["config"] == config]["best_normalized_iteration"]
            * 100
        ).to_numpy(float)
        for config in CONFIG_ORDER
    ]
    _box_strip(
        axes[0], timing, CONFIG_ORDER, "Final public-best first seen (% of budget)"
    )
    depth_groups = [
        cell_summary[
            (cell_summary["config"] == config) & cell_summary["lineage_complete"]
        ]["lineage_depth"].to_numpy(float)
        for config in CONFIG_ORDER
    ]
    depth_labels = [
        f"{config}\nn={len(values)}"
        for config, values in zip(CONFIG_ORDER, depth_groups, strict=True)
    ]
    _box_strip(axes[1], depth_groups, depth_labels, "Final-best lineage depth (edges)")
    fig.suptitle(
        "Figure 14: budget utilization and final-best lineage\n"
        "timing: 70/70 recovered; lineage: 66/70 complete",
        fontsize=11,
        weight="semibold",
    )
    return _save(fig, output_dir, "figure14_budget_lineage")


def figure_cycling(
    cycling: pd.DataFrame, cell_summary: pd.DataFrame, output_dir: Path
) -> list[Path]:
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.5), constrained_layout=True)
    ax = axes[0]
    grid = np.arange(0, 41)
    for base in BASE_ORDER:
        curves = []
        for cell in cell_summary[cell_summary["base_arm"] == base]["cell"]:
            sub = cycling[
                (cycling["cell"] == cell) & cycling["history_complete"]
            ].sort_values("iteration")
            if sub.empty:
                continue
            by_it = dict(
                zip(sub["iteration"], sub["cumulative_recycling_share"], strict=True)
            )
            values, last = [], 0.0
            for iteration in grid:
                if iteration in by_it and pd.notna(by_it[iteration]):
                    last = float(by_it[iteration])
                values.append(last)
            curves.append(values)
        arr = np.asarray(curves, dtype=float)
        median, q1, q3 = (
            np.median(arr, axis=0),
            np.quantile(arr, 0.25, axis=0),
            np.quantile(arr, 0.75, axis=0),
        )
        x = grid / 40
        ax.plot(x, median * 100, label=base, color=COLORS[base], lw=2)
        ax.fill_between(x, q1 * 100, q3 * 100, color=COLORS[base], alpha=0.16)
    ax.set_xlabel("Normalized search iteration")
    ax.set_ylabel("Cumulative recycled added lines (%)")
    ax.set_title("Cumulative cycling over the run")
    ax.legend(frameon=False, fontsize=8)
    _style_axis(ax)

    ax = axes[1]
    configs = CONFIG_ORDER
    bottom = np.zeros(len(configs))
    for column, label, color in (
        ("literal_recycling_share", "Exact line", "#4575b4"),
        ("tuning_recycling_share", "Numeric skeleton", "#fdae61"),
        ("trivial_recycling_share", "Comment/whitespace", "#999999"),
    ):
        values = np.array(
            [
                100 * cell_summary[cell_summary["config"] == config][column].median()
                for config in configs
            ]
        )
        ax.bar(configs, values, bottom=bottom, label=label, color=color, width=0.68)
        bottom += values
    ax.set_ylabel("Median per-cell share of added lines (%)")
    ax.set_title("Cycling type by memory configuration")
    ax.tick_params(axis="x", rotation=35)
    ax.legend(frameon=False, fontsize=7.5)
    _style_axis(ax, "y")
    fig.suptitle(
        "Figure 15: deterministic same-lineage recycling\n"
        "exact and numeric-skeleton matches require a previously deleted ancestor line",
        fontsize=11,
        weight="semibold",
    )
    return _save(fig, output_dir, "figure15_cycling")


def figure_taxonomy(
    taxonomy: pd.DataFrame, edits: pd.DataFrame, output_dir: Path
) -> list[Path]:
    ordered = taxonomy.sort_values("prevalence", ascending=True)
    labels = ordered["display_label"].tolist()
    fig, axes = plt.subplots(1, 3, figsize=(15.0, 5.2), constrained_layout=False)
    fig.subplots_adjust(left=0.09, right=0.985, bottom=0.13, top=0.78, wspace=0.09)
    y = np.arange(len(labels))
    width = 0.23
    for offset, base in zip((-width, 0, width), BASE_ORDER, strict=True):
        key = f"prevalence_{base.lower().replace(' ', '_')}"
        axes[0].barh(
            y + offset, ordered[key] * 100, height=width, label=base, color=COLORS[base]
        )
    axes[0].set_yticks(y, labels)
    axes[0].set_xlabel("Edits carrying label (%)")
    axes[0].set_title("Proxy prevalence by arm")
    axes[0].legend(frameon=False, fontsize=7.5)
    _style_axis(axes[0], "x")

    ratio = ordered["positive_change_odds_ratio"].to_numpy(float)
    low = ordered["odds_ratio_ci_low"].to_numpy(float)
    high = ordered["odds_ratio_ci_high"].to_numpy(float)
    axes[1].errorbar(
        ratio,
        y,
        xerr=np.vstack((ratio - low, high - ratio)),
        fmt="o",
        color="#4c78a8",
        capsize=2.5,
    )
    axes[1].axvline(1, color="#999999", ls="--", lw=1)
    axes[1].set_xscale("log")
    axes[1].set_yticks(y, [])
    axes[1].set_xlabel("Odds ratio for positive score change (log scale)")
    axes[1].set_title("Per-edit helpfulness proxy")
    _style_axis(axes[1], "x")

    h = 0.34
    axes[2].barh(
        y - h / 2,
        ordered["best_update_enrichment"],
        height=h,
        label="Best-so-far updates",
        color="#59a14f",
    )
    axes[2].barh(
        y + h / 2,
        ordered["final_lineage_enrichment"],
        height=h,
        label="Complete final lineages",
        color="#b279a2",
    )
    axes[2].axvline(1, color="#999999", ls="--", lw=1)
    axes[2].set_yticks(y, [])
    axes[2].set_xlabel("Enrichment vs all edits")
    axes[2].set_title("Successful-trajectory enrichment")
    axes[2].legend(frameon=False, fontsize=7.2)
    _style_axis(axes[2], "x")
    fig.suptitle(
        "Figure 16: deterministic edit-taxonomy proxy (screening only)\n"
        "not the paper's LLM-as-judge labels; 95% Wald CI with 0.5 continuity correction",
        fontsize=11,
        weight="semibold",
        y=0.95,
    )
    return _save(fig, output_dir, "figure16_taxonomy_proxy")


def figure_arm_comparison(arms: pd.DataFrame, output_dir: Path) -> list[Path]:
    metrics = [
        ("median_final_loc_ratio", "Final LOC / seed"),
        ("median_final_numeric_literal_ratio", "Final literals / seed"),
        ("median_best_normalized_iteration", "Best first seen / budget"),
        ("median_lineage_depth_complete", "Final lineage depth"),
        ("median_cycling_share", "Cycling share"),
        ("mean_positive_edit_fraction", "Positive-edit fraction"),
    ]
    raw = np.asarray(
        [[float(row[key]) for row in arms.to_dict("records")] for key, _ in metrics]
    )
    standardized = np.zeros_like(raw)
    for i, row in enumerate(raw):
        spread = np.std(row)
        standardized[i] = (row - np.mean(row)) / spread if spread else 0
    fig, ax = plt.subplots(figsize=(10.4, 4.9), constrained_layout=True)
    im = ax.imshow(standardized, aspect="auto", cmap="RdBu_r", vmin=-2, vmax=2)
    ax.set_xticks(range(len(CONFIG_ORDER)), CONFIG_ORDER, rotation=32, ha="right")
    ax.set_yticks(range(len(metrics)), [label for _, label in metrics])
    for i, (key, _) in enumerate(metrics):
        for j, value in enumerate(raw[i]):
            if "iteration" in key or "share" in key or "fraction" in key:
                text = f"{value:.1%}"
            elif "ratio" in key:
                text = f"{value:.2f}×"
            else:
                text = f"{value:.1f}"
            ax.text(j, i, text, ha="center", va="center", fontsize=8, color="black")
    ax.set_title(
        "Figure 17: descriptive memory-arm comparison\n"
        "numbers are raw medians/means; color is standardized within each metric",
        fontsize=11,
        weight="semibold",
    )
    fig.colorbar(im, ax=ax, shrink=0.72, label="Within-metric z-score")
    return _save(fig, output_dir, "figure17_arm_comparison")


def _write_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(path, index=False, quoting=csv.QUOTE_MINIMAL)


def _excel_value(value: Any) -> Any:
    if pd.isna(value) if not isinstance(value, (list, tuple, dict, set)) else False:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def _write_frame(wb: Any, name: str, frame: pd.DataFrame, freeze: str = "A2") -> Any:
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, value in enumerate(frame.columns, start=1):
        cell = ws.cell(1, col, str(value))
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_index, record in enumerate(
        frame.itertuples(index=False, name=None), start=2
    ):
        for col_index, value in enumerate(record, start=1):
            ws.cell(row_index, col_index, _excel_value(value))
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for col_index, column in enumerate(frame.columns, start=1):
        sample = [str(column)] + [
            "" if value is None else str(value)
            for value in frame.iloc[:200, col_index - 1].tolist()
        ]
        ws.column_dimensions[get_column_letter(col_index)].width = min(
            max(len(x) for x in sample) + 2, 42
        )
    return ws


def _key_value_frame(rows: Sequence[tuple[str, Any]]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=["metric", "value"])


def _summary_rows(
    cells: Sequence[dict[str, Any]],
    cell_summary: pd.DataFrame,
    arms: pd.DataFrame,
    taxonomy: pd.DataFrame,
) -> list[tuple[str, Any]]:
    top_prev = taxonomy.sort_values("prevalence", ascending=False).iloc[0]
    top_or = taxonomy.sort_values("positive_change_odds_ratio", ascending=False).iloc[0]
    return [
        ("schema_version", SCHEMA_VERSION),
        ("scope", "70 completed cells: No Memory 10 + GEM 30 + Polyglot 30"),
        ("search_rerun", "No; all values are derived from saved traces/checkpoints"),
        ("tasks", 10),
        ("iterations_per_cell", 40),
        ("observed_parent_child_edits", int(cell_summary["trace_rows"].sum())),
        ("trace_rows_complete_cells", int((cell_summary["trace_rows"] == 40).sum())),
        ("final_best_iteration_coverage", "70/70"),
        (
            "complete_final_lineages",
            f"{int(cell_summary['lineage_complete'].sum())}/70",
        ),
        (
            "median_final_LOC_ratio_all_cells",
            float(cell_summary["final_loc_ratio"].median()),
        ),
        (
            "median_final_numeric_literal_ratio_all_cells",
            float(cell_summary["final_numeric_literal_ratio"].median()),
        ),
        (
            "median_final_best_iteration_fraction",
            float(cell_summary["best_normalized_iteration"].median()),
        ),
        (
            "median_complete_lineage_depth",
            float(cell_summary["lineage_depth"].median()),
        ),
        ("median_cycling_share", float(cell_summary["cycling_share"].median())),
        (
            "cells_with_positive_cycling_slope",
            int((cell_summary["cycling_slope"] > 0).sum()),
        ),
        (
            "most_prevalent_taxonomy_proxy",
            f"{top_prev['display_label']} ({top_prev['prevalence']:.1%})",
        ),
        (
            "largest_helpfulness_OR_proxy",
            f"{top_or['display_label']} (OR={top_or['positive_change_odds_ratio']:.2f})",
        ),
        (
            "taxonomy_warning",
            "Rule-based screening proxy; not the paper's LLM-as-judge labels",
        ),
        (
            "inference_warning",
            "One search run per task/config; comparisons are descriptive, not causal",
        ),
        (
            "trace_warning",
            "JSONL logs accepted/materialized candidates only; missing iterations are not treated as edits",
        ),
        ("paper_method_reference", PAPER),
    ]


def augment_workbook(
    workbook_path: Path,
    output_dir: Path,
    cell_summary: pd.DataFrame,
    shape: pd.DataFrame,
    cycling: pd.DataFrame,
    edits: pd.DataFrame,
    taxonomy: pd.DataFrame,
    arms: pd.DataFrame,
    summary_rows: Sequence[tuple[str, Any]],
    figure_paths: Sequence[Path],
    manifest: Mapping[str, Any],
) -> None:
    wb = load_workbook(workbook_path)
    trace_sheets = [
        "Trace Summary",
        "Shape Trajectory",
        "Budget Lineage",
        "Cycling Detail",
        "Taxonomy Proxy",
        "Taxonomy Summary",
        "Arm Comparison",
        "Trace Visuals",
        "Trace Manifest",
    ]
    for name in trace_sheets:
        if name in wb.sheetnames:
            del wb[name]

    summary_frame = _key_value_frame(summary_rows)
    ws_summary = _write_frame(wb, "Trace Summary", summary_frame)
    ws_summary.column_dimensions["A"].width = 42
    ws_summary.column_dimensions["B"].width = 95
    ws_summary["A1"].fill = PatternFill("solid", fgColor="B4C7E7")
    ws_summary["B1"].fill = PatternFill("solid", fgColor="B4C7E7")

    _write_frame(wb, "Shape Trajectory", shape)
    budget_columns = [
        "cell",
        "task",
        "base_arm",
        "config",
        "trace_rows",
        "iterations_completed",
        "trace_row_coverage",
        "best_iteration",
        "best_normalized_iteration",
        "budget_remaining_fraction",
        "best_iteration_source",
        "lineage_depth",
        "known_lineage_depth",
        "lineage_complete",
        "lineage_error",
        "seed_loc",
        "final_loc",
        "final_loc_ratio",
        "seed_numeric_literals",
        "final_numeric_literals",
        "final_numeric_literal_ratio",
        "cycling_share",
        "cycling_slope",
        "positive_edit_fraction",
        "best_update_count",
    ]
    _write_frame(wb, "Budget Lineage", cell_summary[budget_columns])
    _write_frame(wb, "Cycling Detail", cycling)
    taxonomy_detail_columns = [
        "cell",
        "task",
        "base_arm",
        "config",
        "iteration",
        "parent_id",
        "child_id",
        "parent_score",
        "child_score",
        "score_delta",
        "positive_score_change",
        "best_so_far_update",
        "on_complete_final_lineage",
        "added_nonblank_lines",
        "deleted_nonblank_lines",
        "labels",
        "label_count",
        *LABELS,
    ]
    ws_tax = _write_frame(wb, "Taxonomy Proxy", edits[taxonomy_detail_columns])
    ws_tax.insert_rows(1, amount=3)
    ws_tax["A1"] = "WARNING"
    ws_tax["B1"] = (
        "Deterministic screening proxy; not the paper's validated LLM-as-judge taxonomy."
    )
    ws_tax["A1"].font = Font(bold=True, color="9C0006")
    ws_tax["B1"].font = Font(color="9C0006")
    ws_tax["A2"] = "Summary table"
    ws_tax["B2"] = (
        "See Trace Manifest and trace_analysis_v0.1.0/taxonomy_proxy_summary.csv"
    )
    ws_tax.freeze_panes = "A5"
    _write_frame(wb, "Taxonomy Summary", taxonomy)
    _write_frame(wb, "Arm Comparison", arms)

    ws_visual = wb.create_sheet("Trace Visuals")
    row = 1
    for path in [p for p in figure_paths if p.suffix == ".png"]:
        ws_visual.cell(row, 1, path.stem).font = Font(bold=True, size=12)
        image = XLImage(path)
        image.width = 1000
        image.height = (
            int(image.height * (1000 / image.width)) if image.width else image.height
        )
        # The assignment above changes width before using it; use a fixed readable
        # height instead so Excel anchors never overlap.
        image.height = 460 if "taxonomy" not in path.stem else 500
        ws_visual.add_image(image, f"A{row + 1}")
        row += 29
    ws_visual.sheet_view.showGridLines = False
    ws_visual.column_dimensions["A"].width = 18

    manifest_rows: list[tuple[str, Any]] = []
    for key, value in manifest.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        manifest_rows.append((key, value))
    _write_frame(wb, "Trace Manifest", _key_value_frame(manifest_rows))

    # Add a concise trace section to the workbook's original Summary without
    # changing the private-evaluation section above it.
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        start = 28
        for row in range(start, start + 30):
            for col in range(1, 9):
                ws.cell(row, col).value = None
        ws.cell(start, 1, "Offline trace analyses (all 70 cells)")
        ws.cell(start, 1).font = Font(bold=True, size=12)
        for offset, (key, value) in enumerate(summary_rows[:18], start=1):
            ws.cell(start + offset, 1, key)
            ws.cell(start + offset, 2, _excel_value(value))
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 44)
        ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 90)

    if "Manifest & Sources" in wb.sheetnames:
        ws = wb["Manifest & Sources"]
        rows_to_delete = [
            row
            for row in range(2, ws.max_row + 1)
            if str(ws.cell(row, 1).value or "").startswith("trace_analysis.")
        ]
        for row in reversed(rows_to_delete):
            ws.delete_rows(row)
        for key, value in manifest_rows:
            ws.append((f"trace_analysis.{key}", value))

    # Place new sheets immediately after the original Summary for discoverability.
    ordered = [wb["Summary"]] if "Summary" in wb.sheetnames else []
    ordered += [wb[name] for name in trace_sheets]
    ordered += [ws for ws in wb.worksheets if ws not in ordered]
    wb._sheets = ordered
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    temp = workbook_path.with_suffix(".tmp.xlsx")
    wb.save(temp)
    reopened = load_workbook(temp, read_only=True, data_only=False)
    missing = [name for name in trace_sheets if name not in reopened.sheetnames]
    if missing:
        temp.unlink(missing_ok=True)
        raise ValueError(f"workbook validation failed; missing sheets: {missing}")
    if reopened["Budget Lineage"].max_row != 71:
        temp.unlink(missing_ok=True)
        raise ValueError("workbook validation failed; Budget Lineage is not 70 rows")
    reopened.close()
    os.replace(temp, workbook_path)
    digest = _sha256(workbook_path)
    workbook_path.with_suffix(workbook_path.suffix + ".sha256").write_text(
        f"{digest}  {workbook_path.resolve()}\n", encoding="utf-8"
    )


def render(run_dir: Path, workbook: Path, output_dir: Path) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    cells = load_cells(run_dir)
    shape, cell_summary = shape_trajectories(cells)
    cycling, edits = edit_diagnostics(cells, cell_summary)
    taxonomy = taxonomy_summary(edits)
    arms = arm_summary(cell_summary, edits)

    _write_csv(cell_summary, output_dir / "trace_cell_summary.csv")
    _write_csv(shape, output_dir / "shape_trajectory.csv")
    _write_csv(cycling, output_dir / "cycling_by_edit.csv")
    _write_csv(edits, output_dir / "taxonomy_proxy_by_edit.csv")
    _write_csv(taxonomy, output_dir / "taxonomy_proxy_summary.csv")
    _write_csv(arms, output_dir / "arm_comparison.csv")

    figures = (
        figure_shape(shape, output_dir)
        + figure_budget_lineage(cell_summary, output_dir)
        + figure_cycling(cycling, cell_summary, output_dir)
        + figure_taxonomy(taxonomy, edits, output_dir)
        + figure_arm_comparison(arms, output_dir)
    )
    summary_rows = _summary_rows(cells, cell_summary, arms, taxonomy)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(UTC).isoformat(),
        "run_dir": str(run_dir.resolve()),
        "workbook": str(workbook.resolve()),
        "cells": len(cells),
        "tasks": len({cell["task"] for cell in cells}),
        "config_coverage": dict(Counter(cell["config"] for cell in cells)),
        "observed_trace_edits": len(edits),
        "iterations_completed": int(
            sum(cell["receipt"]["iterations_traced"] for cell in cells)
        ),
        "complete_final_lineages": int(cell_summary["lineage_complete"].sum()),
        "final_best_iteration_coverage": len(cell_summary),
        "cycling_definition": (
            "For each added nonblank physical line, inspect deleted lines in the same "
            "causal ancestor chain. literal = byte-identical; tuning = identical after "
            "numeric literals collapse; trivial = comment/whitespace equivalent."
        ),
        "shape_definition": (
            "Non-comment nonblank LOC and numeric literal occurrences for the public "
            "best-so-far program, normalized to the saved seed. Observed iteration IDs "
            "are retained and gaps are forward-filled."
        ),
        "taxonomy_definition": (
            "Transparent deterministic multi-label rules over code diffs and saved change "
            "descriptions. Screening proxy only; not the paper's LLM-as-judge classifier."
        ),
        "paper_method_reference": PAPER,
        "limitations": [
            "40 iterations per cell versus 100 in the paper",
            "one Qwen3.8/OpenEvolve-style setup versus four frameworks and five models",
            "one search seed per task/config; configuration comparisons are descriptive",
            "JSONL contains materialized candidates, not one row for every completed iteration",
            "4 of 70 final-best causal lineages have a missing ancestor checkpoint",
            "taxonomy proxy has no human reliability validation and is not paper-comparable",
        ],
        "outputs": sorted(path.name for path in output_dir.iterdir() if path.is_file()),
    }
    manifest_path = output_dir / "trace_analysis_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    manifest["outputs"] = sorted(
        path.name for path in output_dir.iterdir() if path.is_file()
    )
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    augment_workbook(
        workbook,
        output_dir,
        cell_summary,
        shape,
        cycling,
        edits,
        taxonomy,
        arms,
        summary_rows,
        figures,
        manifest,
    )
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    manifest = render(args.run_dir, args.workbook, args.output_dir)
    print(
        json.dumps(
            {
                "cells": manifest["cells"],
                "observed_trace_edits": manifest["observed_trace_edits"],
                "complete_final_lineages": manifest["complete_final_lineages"],
                "workbook": manifest["workbook"],
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
