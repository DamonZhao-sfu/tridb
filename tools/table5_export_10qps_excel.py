#!/usr/bin/env python3
"""Export the accepted Track C 10-QPS Search/Add receipts to one workbook."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
RESULT_ROOT = REPO / "bench/out/table5_reproduction_2026_08_19"
SWEEP_ROOT = REPO / "bench/out/table5_track_c_qps_sweep_2026_08_20_v1"
OUTPUT = REPO / "haikaidocs/table5_track_c_10qps_search_add_comparison_2026-08-20.xlsx"

DISPLAY = {
    "tridb_gem": "TriDB/GEM",
    "mem0": "Mem0 2.0.18",
    "cognee": "Cognee 1.5.0",
    "memos": "MemOS 2.0.30",
}

CANONICAL = {
    ("tridb_gem", "search"): RESULT_ROOT / "runs/tridb_gem/search/b1/run_receipt.json",
    ("mem0", "search"): RESULT_ROOT / "runs/mem0/search/b1/run_receipt.json",
    ("cognee", "search"): RESULT_ROOT / "runs/cognee/search/b1/run_receipt.json",
    ("memos", "search"): RESULT_ROOT / "runs/memos/search/b1/run_receipt.json",
    ("tridb_gem", "add:native"): RESULT_ROOT
    / "concurrent_gpu1_pilot_v1/runs/tridb_gem/add_native/b1/run_receipt.json",
    ("mem0", "add:native"): RESULT_ROOT
    / "concurrent_gpu1_pilot_v1/runs/mem0/add_native/b1/run_receipt.json",
    ("cognee", "add:native"): RESULT_ROOT
    / "concurrent_gpu1_pilot_v1/runs/cognee/add_native/b1/run_receipt.json",
    ("memos", "add:native"): RESULT_ROOT
    / "concurrent_gpu1_pilot_v2_retry1/runs/memos/add_native/b1/run_receipt.json",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def timeout_count(receipt: dict[str, Any]) -> int:
    errors = receipt.get("formal_summary", {}).get("errors", {})
    return sum(int(count) for error, count in errors.items() if "TimeoutError" in error)


def seconds(value: Any) -> float | None:
    return None if value is None else float(value) / 1000.0


def latency(summary: dict[str, Any], family: str, key: str) -> float | None:
    return seconds((summary.get(family) or {}).get(key))


def configure_sheet(ws: Any, *, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for column in range(1, ws.max_column + 1):
        width = min(
            70,
            max(
                11,
                max(
                    len(str(ws.cell(row=row, column=column).value or ""))
                    for row in range(1, ws.max_row + 1)
                )
                + 2,
            ),
        )
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_table(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    configure_sheet(ws)


def validate_canonical(receipts: dict[tuple[str, str], dict[str, Any]]) -> None:
    if set(receipts) != set(CANONICAL):
        raise RuntimeError("canonical receipt coverage is incomplete")
    for key, receipt in receipts.items():
        protocol = receipt.get("protocol", {})
        if receipt.get("status") != "complete":
            raise RuntimeError(f"canonical receipt is incomplete: {key}")
        if float(protocol.get("qps")) != 10.0:
            raise RuntimeError(f"canonical receipt is not 10 QPS: {key}")
        if float(protocol.get("timeout_seconds")) != 60.0:
            raise RuntimeError(f"canonical receipt is not timeout=60s: {key}")


def main() -> None:
    receipts = {key: load(path) for key, path in CANONICAL.items()}
    validate_canonical(receipts)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Track C 10-QPS Search and Native Add comparison"
    workbook.properties.creator = "TriDB Track C benchmark harness"
    workbook.properties.description = (
        "Uniform 10-QPS, 60-second-deadline comparison with source receipt paths"
    )

    systems = ("tridb_gem", "mem0", "cognee", "memos")
    summary_rows = []
    for system in systems:
        search = receipts[(system, "search")]
        add = receipts[(system, "add:native")]
        ss = search["formal_summary"]
        ads = add["formal_summary"]
        quality = search.get("quality", {})
        summary_rows.append(
            [
                DISPLAY[system],
                ss["successful"],
                ss["total"],
                ss["success_rate"],
                timeout_count(search),
                latency(ss, "service_latency_ms", "mean"),
                latency(ss, "user_visible_latency_ms", "p99"),
                ss.get("completion_throughput_qps") or ss.get("actual_qps"),
                quality.get("evidence_recall_at_20"),
                quality.get("evidence_any_hit_at_20"),
                ads["successful"],
                ads["total"],
                ads["success_rate"],
                timeout_count(add),
                latency(ads, "service_latency_ms", "mean"),
                latency(ads, "user_visible_latency_ms", "p99"),
                ads.get("completion_throughput_qps") or ads.get("actual_qps"),
            ]
        )
    ws = workbook.create_sheet("10QPS Summary")
    write_table(
        ws,
        [
            "System",
            "Search success",
            "Search total",
            "Search success rate",
            "Search timeouts",
            "Search service mean (s, successes)",
            "Search user-visible P99 (s, all)",
            "Search completion QPS",
            "Search effective Recall@20",
            "Search effective Any-hit@20",
            "Add success",
            "Add total",
            "Add success rate",
            "Add timeouts",
            "Add service mean (s, successes)",
            "Add user-visible P99 (s, all)",
            "Add completion QPS",
        ],
        summary_rows,
    )
    for row in range(2, ws.max_row + 1):
        for column in (4, 9, 10, 13):
            ws.cell(row, column).number_format = "0.00%"
        for column in (6, 7, 15, 16):
            ws.cell(row, column).number_format = "0.000"
        fill = GOOD_FILL if ws.cell(row, 4).value == 1 else WARN_FILL
        ws.cell(row, 4).fill = fill
        ws.cell(row, 13).fill = GOOD_FILL if ws.cell(row, 13).value == 1 else WARN_FILL

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Success rate at 10 QPS"
    chart.y_axis.title = "Success rate"
    chart.x_axis.title = "System"
    chart.height = 8
    chart.width = 15
    chart.add_data(
        Reference(ws, min_col=4, max_col=4, min_row=1, max_row=5), titles_from_data=True
    )
    chart.add_data(
        Reference(ws, min_col=13, max_col=13, min_row=1, max_row=5),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(chart, "A8")

    def detail_rows(phase: str) -> list[list[Any]]:
        rows = []
        for system in systems:
            receipt = receipts[(system, phase)]
            formal = receipt["formal_summary"]
            quality = receipt.get("quality", {})
            rows.append(
                [
                    DISPLAY[system],
                    formal["total"],
                    formal["successful"],
                    formal["failed"],
                    timeout_count(receipt),
                    formal["success_rate"],
                    formal.get("admission_qps"),
                    formal.get("completion_throughput_qps") or formal.get("actual_qps"),
                    latency(formal, "service_latency_ms", "mean"),
                    latency(formal, "service_latency_ms", "p50"),
                    latency(formal, "service_latency_ms", "p90"),
                    latency(formal, "service_latency_ms", "p95"),
                    latency(formal, "service_latency_ms", "p99"),
                    latency(formal, "user_visible_latency_ms", "mean"),
                    latency(formal, "user_visible_latency_ms", "p50"),
                    latency(formal, "user_visible_latency_ms", "p90"),
                    latency(formal, "user_visible_latency_ms", "p95"),
                    latency(formal, "user_visible_latency_ms", "p99"),
                    quality.get("evidence_recall_at_20"),
                    quality.get("evidence_any_hit_at_20"),
                    receipt["build_id"],
                    str(CANONICAL[(system, phase)].resolve()),
                ]
            )
        return rows

    detail_headers = [
        "System",
        "Total",
        "Successful",
        "Failed",
        "Timeouts",
        "Success rate",
        "Admission QPS",
        "Completion QPS",
        "Service mean (s, successes)",
        "Service P50 (s, successes)",
        "Service P90 (s, successes)",
        "Service P95 (s, successes)",
        "Service P99 (s, successes)",
        "User-visible mean (s, all)",
        "User-visible P50 (s, all)",
        "User-visible P90 (s, all)",
        "User-visible P95 (s, all)",
        "User-visible P99 (s, all)",
        "Effective Recall@20",
        "Effective Any-hit@20",
        "Build ID",
        "Receipt path",
    ]
    for title, phase in (
        ("Search 10QPS", "search"),
        ("Native Add 10QPS", "add:native"),
    ):
        ws_detail = workbook.create_sheet(title)
        write_table(ws_detail, detail_headers, detail_rows(phase))
        for row in range(2, ws_detail.max_row + 1):
            ws_detail.cell(row, 6).number_format = "0.00%"
            for column in range(9, 19):
                ws_detail.cell(row, column).number_format = "0.000"
            for column in (19, 20):
                ws_detail.cell(row, column).number_format = "0.00%"

    all_rows = []
    canonical_paths = {path.resolve() for path in CANONICAL.values()}
    for path in sorted(RESULT_ROOT.glob("**/run_receipt.json")):
        receipt = load(path)
        if (
            receipt.get("status") != "complete"
            or receipt.get("protocol", {}).get("qps") != 10.0
        ):
            continue
        formal = receipt.get("formal_summary", {})
        quality = receipt.get("quality", {})
        all_rows.append(
            [
                path.resolve() in canonical_paths,
                DISPLAY.get(receipt.get("system"), receipt.get("system")),
                receipt.get("phase"),
                receipt.get("build_id"),
                receipt.get("protocol", {}).get("timeout_seconds"),
                formal.get("total"),
                formal.get("successful"),
                formal.get("failed"),
                timeout_count(receipt),
                formal.get("success_rate"),
                latency(formal, "service_latency_ms", "mean"),
                latency(formal, "user_visible_latency_ms", "p99"),
                formal.get("completion_throughput_qps") or formal.get("actual_qps"),
                quality.get("evidence_recall_at_20"),
                str(path.resolve()),
            ]
        )
    ws_all = workbook.create_sheet("All 10QPS Receipts")
    write_table(
        ws_all,
        [
            "Canonical main comparison",
            "System",
            "Phase",
            "Build ID",
            "Timeout (s)",
            "Total",
            "Successful",
            "Failed",
            "Timeouts",
            "Success rate",
            "Service mean (s, successes)",
            "User-visible P99 (s, all)",
            "Completion QPS",
            "Effective Recall@20",
            "Receipt path",
        ],
        all_rows,
    )
    for row in range(2, ws_all.max_row + 1):
        ws_all.cell(row, 10).number_format = "0.00%"
        ws_all.cell(row, 14).number_format = "0.00%"
        if ws_all.cell(row, 1).value:
            for cell in ws_all[row]:
                cell.fill = GOOD_FILL

    progress_rows: list[list[Any]] = []
    sweep_receipt_path = SWEEP_ROOT / "sweep_receipt.json"
    if sweep_receipt_path.exists():
        sweep = load(sweep_receipt_path)
        for point in sweep.get("points", []):
            point_receipt = Path(point["run_dir"]) / "run_receipt.json"
            native = load(point_receipt) if point_receipt.exists() else {}
            formal = native.get("formal_summary", {})
            progress_rows.append(
                [
                    DISPLAY.get(point["system"], point["system"]),
                    point["workload"],
                    point["qps"],
                    point["status"],
                    native.get("status"),
                    formal.get("total"),
                    formal.get("successful"),
                    point["run_dir"],
                ]
            )
    ws_progress = workbook.create_sheet("QPS 1-5 Progress")
    write_table(
        ws_progress,
        [
            "System",
            "Workload",
            "QPS",
            "Sweep status",
            "Native receipt status",
            "Formal total",
            "Formal successful",
            "Run directory",
        ],
        progress_rows,
    )

    ws_notes = workbook.create_sheet("Definitions and Notes")
    notes = [
        ("Generated at", datetime.now(timezone.utc).isoformat()),
        (
            "Main comparison",
            "Only canonical 10-QPS receipts with the common 60-second deadline.",
        ),
        ("Search workload", "LoCoMo Search B1, 1,787 formal queries, top-k=35."),
        (
            "Add workload",
            "Native Add, 2,000 formal admissions; Source-to-searchable is excluded.",
        ),
        (
            "Service latency",
            "Measured from adapter call start to completion; distribution contains successful requests only.",
        ),
        (
            "User-visible latency",
            "Scheduled admission to client-visible completion/timeout; distribution contains all requests.",
        ),
        (
            "Success rate",
            "Successful formal requests divided by all formal admissions.",
        ),
        (
            "Effective Recall@20",
            "Evidence Recall@20 with timeout/failure requests retained in the denominator and scored as zero.",
        ),
        (
            "Completion QPS",
            "Observed completions divided by the formal run wall interval; overloaded systems may be below admission QPS.",
        ),
        (
            "Cognee Search",
            "Main sheet uses the common timeout=60s run. The timeout=300s sensitivity receipt remains visible only in All 10QPS Receipts.",
        ),
        (
            "Cognee Native Add",
            "Measures native source persistence only; cognify/source-to-searchable work is not included.",
        ),
        (
            "Repeat limitation",
            "These are accepted single-run B1/pilot receipts, not three-repeat confidence intervals.",
        ),
        (
            "Hardware note",
            "Search and Add accepted receipts used the same controlled models but were collected in the recorded GPU0/GPU1 execution slots; consult receipt paths for provenance.",
        ),
    ]
    write_table(
        ws_notes,
        ["Item", "Definition / caveat"],
        [[key, value] for key, value in notes],
    )
    ws_notes.column_dimensions["A"].width = 28
    ws_notes.column_dimensions["B"].width = 110

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
