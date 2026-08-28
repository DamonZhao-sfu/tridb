#!/usr/bin/env python3
"""Audit and export the Track C QPS sweep without mutating it."""

from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPECTED_DATASET_SHA256 = (
    "79fa87e90f04081343b8c8debecb80a9a6842b76a7aa537dc9fdf651ea698ff4"
)
SYSTEMS = (
    "tridb_gem",
    "mem0",
    "cognee",
    "memos",
    "mandol",
    "graphiti_zep_oss_proxy",
)
SYSTEM_LABELS = {
    "tridb_gem": "TriDB/GEM",
    "mem0": "Mem0 2.0.18",
    "cognee": "Cognee 1.5.0",
    "memos": "MemOS 2.0.30",
    "mandol": "Mandol 0.1.0",
    "graphiti_zep_oss_proxy": "Graphiti (Zep OSS proxy)",
}
WORKLOADS = ("search", "add_native")
QPS_VALUES = (1, 5, 10)


def _scan_jsonl(path: Path) -> dict[str, Any]:
    count = 0
    parse_errors = 0
    indices: set[int] = set()
    trace_ids: set[str] = set()
    phases: set[str] = set()
    systems: set[str] = set()
    build_ids: set[str] = set()
    if not path.is_file():
        return {
            "exists": False,
            "count": 0,
            "parse_errors": 0,
            "indices": [],
            "trace_ids": [],
            "phases": [],
            "systems": [],
            "build_ids": [],
        }
    with path.open(encoding="utf-8") as source:
        for line in source:
            if not line.strip():
                continue
            count += 1
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                parse_errors += 1
                continue
            index = record.get("request_index")
            if isinstance(index, int) and not isinstance(index, bool):
                indices.add(index)
            trace_id = record.get("trace_id")
            if isinstance(trace_id, str) and trace_id:
                trace_ids.add(trace_id)
            for key, target in (
                ("phase", phases),
                ("system", systems),
                ("build_id", build_ids),
            ):
                value = record.get(key)
                if isinstance(value, str):
                    target.add(value)
    return {
        "exists": True,
        "count": count,
        "parse_errors": parse_errors,
        "indices": sorted(indices),
        "trace_ids": sorted(trace_ids),
        "phases": sorted(phases),
        "systems": sorted(systems),
        "build_ids": sorted(build_ids),
    }


def _scan_spans(path: Path, formal_phase: str, warmup_phase: str) -> dict[str, Any]:
    count = 0
    parse_errors = 0
    formal_roots: set[str] = set()
    warmup_roots: set[str] = set()
    categories: set[str] = set()
    if not path.is_file():
        return {
            "exists": False,
            "count": 0,
            "parse_errors": 0,
            "formal_root_trace_ids": [],
            "warmup_root_trace_ids": [],
            "categories": [],
        }
    with path.open(encoding="utf-8") as source:
        for line in source:
            if not line.strip():
                continue
            count += 1
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                parse_errors += 1
                continue
            category = record.get("category")
            if isinstance(category, str):
                categories.add(category)
            if category != "request" or record.get("parent_span_id") is not None:
                continue
            trace_id = record.get("trace_id")
            if not isinstance(trace_id, str) or not trace_id:
                continue
            if record.get("phase") == formal_phase:
                formal_roots.add(trace_id)
            elif record.get("phase") == warmup_phase:
                warmup_roots.add(trace_id)
    return {
        "exists": True,
        "count": count,
        "parse_errors": parse_errors,
        "formal_root_trace_ids": sorted(formal_roots),
        "warmup_root_trace_ids": sorted(warmup_roots),
        "categories": sorted(categories),
    }


def _timeout_count(receipt: dict[str, Any]) -> int | None:
    summary = receipt.get("formal_summary")
    if not isinstance(summary, dict):
        return None
    return sum(
        int(count)
        for error, count in (summary.get("errors") or {}).items()
        if "TimeoutError" in str(error)
    )


def _nested(payload: dict[str, Any], *keys: str) -> Any:
    value: Any = payload
    for key in keys:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def _orchestrator_point(
    root: Path, run_dir: Path, system: str, workload: str, qps: int
) -> dict[str, Any] | None:
    """Find a root-level sweep point that owns this exact run directory."""
    resolved_run = str(run_dir.resolve())
    for path in sorted(root.glob("*sweep*receipt*.json")):
        try:
            receipt = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        for point in receipt.get("points") or []:
            if (
                point.get("system") == system
                and point.get("workload") == workload
                and float(point.get("qps", -1)) == float(qps)
                and str(Path(point.get("run_dir", "")).resolve()) == resolved_run
            ):
                return {
                    "receipt": str(path.resolve()),
                    "sweep_status": receipt.get("status"),
                    "point_status": point.get("status"),
                }
    return None


def audit_point(root: Path, system: str, workload: str, qps: int) -> dict[str, Any]:
    run_dir = root / "runs" / system / workload / f"qps_{qps}"
    receipt_path = run_dir / "run_receipt.json"
    expected_formal = 1_787 if workload == "search" else 2_000
    expected_warmup = 307 if workload == "search" else 10
    formal_phase = "formal_search" if workload == "search" else "formal_add_native"
    warmup_phase = "warmup_search" if workload == "search" else "warmup_add_native"
    if not receipt_path.is_file():
        return {
            "system": system,
            "workload": workload,
            "qps": qps,
            "status": "MISSING",
            "valid": False,
            "receipt_path": str(receipt_path.resolve()),
            "checks": {},
        }

    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    orchestrator = _orchestrator_point(root, run_dir, system, workload, qps)
    formal = _scan_jsonl(run_dir / "formal.jsonl")
    warmup = _scan_jsonl(run_dir / "warmup.jsonl")
    spans = _scan_spans(run_dir / "spans.jsonl", formal_phase, warmup_phase)
    summary = receipt.get("formal_summary") or {}
    checks = {
        "receipt_complete": receipt.get("status") == "complete",
        "dataset_checksum": _nested(receipt, "dataset", "sha256")
        == EXPECTED_DATASET_SHA256,
        "qps_exact": _nested(receipt, "protocol", "qps") == float(qps),
        "timeout_60s": _nested(receipt, "protocol", "timeout_seconds") == 60.0,
        "stage_profiling": _nested(receipt, "protocol", "stage_profiling") is True,
        "formal_count": formal["count"] == expected_formal,
        "formal_parse": formal["parse_errors"] == 0,
        "formal_indices": formal["indices"] == list(range(expected_formal)),
        "formal_trace_ids": len(formal["trace_ids"]) == expected_formal,
        "formal_phase": formal["phases"] == [formal_phase],
        "formal_system": formal["systems"] == [system],
        "warmup_count": warmup["count"] == expected_warmup,
        "warmup_parse": warmup["parse_errors"] == 0,
        "warmup_indices": warmup["indices"] == list(range(expected_warmup)),
        "warmup_trace_ids": len(warmup["trace_ids"]) == expected_warmup,
        "warmup_phase": warmup["phases"] == [warmup_phase],
        "span_parse": spans["parse_errors"] == 0,
        "formal_root_span_per_request": spans["formal_root_trace_ids"]
        == formal["trace_ids"],
        "warmup_root_span_per_request": spans["warmup_root_trace_ids"]
        == warmup["trace_ids"],
        "summary_total": summary.get("total") == expected_formal,
        "time_breakdown": isinstance(receipt.get("time_breakdown"), dict),
        "orchestrator_point_complete": orchestrator is None
        or orchestrator.get("point_status") == "complete",
    }
    valid = all(checks.values())
    formal_valid = all(
        passed
        for name, passed in checks.items()
        if name != "orchestrator_point_complete"
    )
    status = (
        "COMPLETE_VALID"
        if valid
        else "FORMAL_VALID_FINALIZING"
        if formal_valid
        and orchestrator is not None
        and orchestrator.get("point_status") == "running"
        else "RUNNING"
        if receipt.get("status") == "running"
        else "COMPLETE_INVALID"
        if receipt.get("status") == "complete"
        else str(receipt.get("status") or "UNKNOWN").upper()
    )
    return {
        "system": system,
        "workload": workload,
        "qps": qps,
        "status": status,
        "valid": valid,
        "formal_valid": formal_valid,
        "orchestrator": orchestrator,
        "receipt_status": receipt.get("status"),
        "receipt_path": str(receipt_path.resolve()),
        "build_id": receipt.get("build_id"),
        "formal_records": formal["count"],
        "warmup_records": warmup["count"],
        "span_records": spans["count"],
        "successful": summary.get("successful"),
        "failed": summary.get("failed"),
        "timeouts": _timeout_count(receipt),
        "success_rate": summary.get("success_rate"),
        "service_mean_ms": _nested(summary, "service_latency_ms", "mean"),
        "service_p90_ms": _nested(summary, "service_latency_ms", "p90"),
        "service_p99_ms": _nested(summary, "service_latency_ms", "p99"),
        "user_visible_p99_ms": _nested(summary, "user_visible_latency_ms", "p99"),
        "completion_qps": summary.get("completion_throughput_qps")
        or summary.get("actual_qps"),
        "evidence_recall_at_20": _nested(receipt, "quality", "evidence_recall_at_20"),
        "evidence_any_hit_at_20": _nested(receipt, "quality", "evidence_any_hit_at_20"),
        "checks": checks,
    }


def _configure_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(70, width + 2)


def _write_workbook(path: Path, points: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "System",
        "Workload",
        "QPS",
        "Audit status",
        "Formal",
        "Successful",
        "Failed",
        "Timeouts",
        "Success rate",
        "Service mean (ms, success-only)",
        "Service P90 (ms, success-only)",
        "Service P99 (ms, success-only)",
        "User-visible P99 (ms, all)",
        "Completion QPS",
        "Evidence Recall@20",
        "Evidence any-hit@20",
        "Spans",
        "Receipt",
        "Orchestrator status",
        "Orchestrator receipt",
    ]
    for workload, title in (("search", "Search"), ("add_native", "Native Add")):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for point in points:
            if point["workload"] != workload:
                continue
            sheet.append(
                [
                    SYSTEM_LABELS[point["system"]],
                    workload,
                    point["qps"],
                    point["status"],
                    point.get("formal_records"),
                    point.get("successful"),
                    point.get("failed"),
                    point.get("timeouts"),
                    point.get("success_rate"),
                    point.get("service_mean_ms"),
                    point.get("service_p90_ms"),
                    point.get("service_p99_ms"),
                    point.get("user_visible_p99_ms"),
                    point.get("completion_qps"),
                    point.get("evidence_recall_at_20"),
                    point.get("evidence_any_hit_at_20"),
                    point.get("span_records"),
                    point["receipt_path"],
                    _nested(point, "orchestrator", "point_status"),
                    _nested(point, "orchestrator", "receipt"),
                ]
            )
        _configure_sheet(sheet)
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, 9).number_format = "0.00%"
            sheet.cell(row, 15).number_format = "0.00%"
            sheet.cell(row, 16).number_format = "0.00%"
            status = sheet.cell(row, 4).value
            color = "E2F0D9" if status == "COMPLETE_VALID" else "FFF2CC"
            sheet.cell(row, 4).fill = PatternFill("solid", fgColor=color)

    audit_sheet = workbook.create_sheet("Trace Audit")
    check_names = sorted({name for point in points for name in point.get("checks", {})})
    audit_sheet.append(["System", "Workload", "QPS", "Status", *check_names])
    for point in points:
        audit_sheet.append(
            [
                SYSTEM_LABELS[point["system"]],
                point["workload"],
                point["qps"],
                point["status"],
                *[point.get("checks", {}).get(name) for name in check_names],
            ]
        )
    _configure_sheet(audit_sheet)
    workbook.save(path)


def _write_markdown(
    path: Path, points: list[dict[str, Any]], generated_at: str
) -> None:
    lines = [
        "# Track C QPS sweep audit",
        "",
        f"Generated: `{generated_at}`",
        "",
        "`COMPLETE_VALID` means the receipt, exact formal/warmup cardinalities, "
        "request indices, trace IDs, root request spans, dataset checksum, QPS, "
        "timeout, and time-breakdown gates all passed.",
        "",
        "| System | Workload | QPS | Status | Success/total | Success rate | "
        "Mean ms | P90 ms | P99 ms | Completion QPS |",
        "|---|---|---:|---|---:|---:|---:|---:|---:|---:|",
    ]
    for point in points:
        successful = point.get("successful")
        total = point.get("formal_records")
        ratio = "—" if successful is None else f"{successful}/{total}"
        rate = point.get("success_rate")
        rate_text = "—" if rate is None else f"{100 * float(rate):.2f}%"

        def number(key: str) -> str:
            value = point.get(key)
            return "—" if value is None else f"{float(value):.3f}"

        lines.append(
            f"| {SYSTEM_LABELS[point['system']]} | {point['workload']} | "
            f"{point['qps']} | {point['status']} | {ratio} | {rate_text} | "
            f"{number('service_mean_ms')} | {number('service_p90_ms')} | "
            f"{number('service_p99_ms')} | {number('completion_qps')} |"
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--require-complete", action="store_true")
    args = parser.parse_args()
    root = Path(args.root).resolve()
    output = Path(args.output_dir).resolve()
    if output.exists():
        raise FileExistsError(f"refusing existing audit output: {output}")
    output.mkdir(parents=True)
    generated_at = datetime.now(timezone.utc).isoformat()
    points = [
        audit_point(root, system, workload, qps)
        for workload in WORKLOADS
        for system in SYSTEMS
        for qps in QPS_VALUES
    ]
    payload = {
        "schema_version": "table5_track_c_qps_audit_v0.1.0",
        "generated_at": generated_at,
        "source_root": str(root),
        "expected_dataset_sha256": EXPECTED_DATASET_SHA256,
        "complete_valid_points": sum(point["valid"] for point in points),
        "expected_points": len(points),
        "all_points_valid": all(point["valid"] for point in points),
        "points": points,
    }
    (output / "qps_audit.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_markdown(output / "QPS_AUDIT.md", points, generated_at)
    _write_workbook(output / "qps_comparison.xlsx", points)
    print(
        json.dumps(
            {
                "output": str(output),
                "complete_valid_points": payload["complete_valid_points"],
                "expected_points": payload["expected_points"],
                "all_points_valid": payload["all_points_valid"],
            }
        )
    )
    return int(args.require_complete and not payload["all_points_valid"])


if __name__ == "__main__":
    raise SystemExit(main())
