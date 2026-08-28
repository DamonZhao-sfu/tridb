"""Export EvoMemBench Table 3/5 and three-system latency artifacts.

The workbook deliberately separates paper reference values from locally
measured values.  The local CrossEp-Know run is complete, whereas an INEP-KNOW
Table-3-shaped run must not be represented as complete until all six official
headline subsets have answer predictions under the declared context policy.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import math
from pathlib import Path
from statistics import mean, median
from typing import Any, Iterable, Sequence


PAPER_TABLE3 = (
    (
        "Long-Context LLM",
        "Gemini-3-Flash",
        98.20,
        83.00,
        94.00,
        77.00,
        1.00,
        58.00,
        96.00,
        1.00,
        1.00,
    ),
    (
        "Long-Context LLM",
        "GPT-5-mini",
        80.60,
        62.33,
        85.00,
        73.00,
        3.50,
        24.00,
        77.00,
        2.00,
        3.00,
    ),
    (
        "Long-Context LLM",
        "DeepSeek-V3.2",
        83.00,
        32.33,
        64.00,
        42.00,
        6.75,
        10.00,
        65.00,
        3.00,
        5.50,
    ),
    (
        "Retrieval-Augmented Memory",
        "BM25",
        88.80,
        50.33,
        73.00,
        55.00,
        4.25,
        7.00,
        54.00,
        4.00,
        4.17,
    ),
    (
        "Retrieval-Augmented Memory",
        "Qwen3-Emb-4B",
        83.60,
        21.33,
        38.00,
        38.00,
        8.50,
        4.00,
        30.00,
        7.50,
        8.17,
    ),
    (
        "Retrieval-Augmented Memory",
        "GraphRAG",
        74.20,
        22.00,
        31.00,
        31.00,
        10.50,
        4.00,
        18.00,
        9.50,
        10.17,
    ),
    (
        "Short-Term Memory",
        "MemAgent",
        71.80,
        12.67,
        26.00,
        39.00,
        10.75,
        3.00,
        19.00,
        10.00,
        10.50,
    ),
    (
        "Short-Term Memory",
        "MemoBrain",
        77.60,
        10.67,
        30.00,
        32.00,
        11.00,
        1.00,
        20.00,
        11.00,
        11.00,
    ),
    (
        "General Long-Term Memory",
        "Mem0",
        79.20,
        52.00,
        75.00,
        68.00,
        4.50,
        7.00,
        49.00,
        4.50,
        4.50,
    ),
    (
        "General Long-Term Memory",
        "A-MEM",
        91.20,
        51.00,
        60.00,
        46.00,
        4.25,
        5.00,
        35.00,
        6.00,
        4.83,
    ),
    (
        "General Long-Term Memory",
        "MemOS",
        90.40,
        41.33,
        55.00,
        43.00,
        5.50,
        3.00,
        35.00,
        7.50,
        6.17,
    ),
    (
        "General Long-Term Memory",
        "MemoryOS",
        84.80,
        38.67,
        41.00,
        37.00,
        7.50,
        2.00,
        30.00,
        9.50,
        8.17,
    ),
)

PAPER_TABLE3_HEADERS = (
    "Category",
    "Method",
    "EventQA (%)",
    "LME (S*) (%)",
    "Ruler qa1 (%)",
    "Ruler qa2 (%)",
    "Retention Avg Rank",
    "FC MH (%)",
    "FC SH (%)",
    "Revision Avg Rank",
    "Overall Rank",
)

INEP_COVERAGE = (
    (
        "EventQA",
        "Memory Retention",
        500,
        "eventqa_full",
        "binary eventqa_recall",
        411_984,
        752_514,
        False,
    ),
    (
        "LME (S*)",
        "Memory Retention",
        300,
        "longmemeval_s*",
        "official LLM yes/no judge",
        359_131,
        392_274,
        False,
    ),
    (
        "Ruler qa1",
        "Memory Retention",
        100,
        "ruler_qa1_197K",
        "substring exact match",
        212_966,
        212_966,
        True,
    ),
    (
        "Ruler qa2",
        "Memory Retention",
        100,
        "ruler_qa2_421K",
        "substring exact match",
        467_102,
        467_102,
        False,
    ),
    (
        "FC MH",
        "Memory Revision",
        100,
        "factconsolidation_mh_262k",
        "substring exact match",
        332_863,
        332_863,
        False,
    ),
    (
        "FC SH",
        "Memory Revision",
        100,
        "factconsolidation_sh_262k",
        "substring exact match",
        332_863,
        332_863,
        False,
    ),
)

ARM_LABELS = {"memory_off": "No Memory", "full_gem": "GEM", "polyglot": "Polyglot"}
ARM_COLORS = {"No Memory": "#4C78A8", "GEM": "#F58518", "Polyglot": "#54A24B"}


def _json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text())


def _jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def _percentile(values: Sequence[float], fraction: float) -> float:
    ordered = sorted(float(value) for value in values)
    return ordered[max(0, math.ceil(fraction * len(ordered)) - 1)]


def _dist(values: Iterable[float]) -> dict[str, float | int]:
    materialized = list(values)
    return {
        "n": len(materialized),
        "mean": mean(materialized),
        "p50": median(materialized),
        "p95": _percentile(materialized, 0.95),
        "p99": _percentile(materialized, 0.99),
    }


def _trace_map(
    paths: Sequence[Path], arm: str | None = None
) -> dict[str, dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in paths:
        rows.extend(_jsonl(path))
    if arm is not None:
        rows = [
            row for row in rows if row["arm"] == arm and int(row["history_size"]) > 0
        ]
    mapped = {str(row["target_id"]): row for row in rows}
    if len(mapped) != len(rows):
        raise ValueError("duplicate trace target IDs")
    return mapped


def _phase_rows(
    run_root: Path, poly_root: Path
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    no = _trace_map(
        [run_root / f"shard_{i}/traces/memory_off.jsonl" for i in (0, 1)], "memory_off"
    )
    gem = _trace_map(
        [run_root / f"shard_{i}/traces/full_gem.jsonl" for i in (0, 1)], "full_gem"
    )
    poly = _trace_map(
        [poly_root / f"polyglot_{i}/traces/multi_system.jsonl" for i in (0, 1)]
    )
    if not (set(no) == set(gem) == set(poly)) or len(gem) != 764:
        raise ValueError("expected exactly 764 paired three-system reuse traces")

    per_query: list[dict[str, Any]] = []
    for target in sorted(gem):
        nlat = no[target]["latency_ms"]
        glat = gem[target]["latency_ms"]
        plat = poly[target]["latency_ms"]
        systems = {
            "No Memory": {
                "query_embedding": 0.0,
                "vector_search": 0.0,
                "relation_scan": 0.0,
                "graph_traversal": 0.0,
                "merge_memory": 0.0,
                "prompt_processing": float(nlat["prompt_serialization"])
                + float(nlat["prompt_tokenization"]),
                "model_prefill": float(nlat["model_prefill"]),
                "model_decode": float(nlat["model_decode"]),
                "end_to_end": float(nlat["end_to_end"]),
                "database_retrieve": None,
                "effective_memory_path": None,
                "e2e_provenance": "measured",
            },
            "GEM": {
                "query_embedding": float(glat["query_embedding"]),
                # GEM exposes fused DB work as one streaming operator.  It is not
                # valid to fabricate isolated vector/graph/filter wall times.
                "vector_search": 0.0,
                "relation_scan": 0.0,
                "graph_traversal": 0.0,
                "merge_memory": float(glat["database_retrieval"]),
                "prompt_processing": float(glat["prompt_serialization"])
                + float(glat["prompt_tokenization"]),
                "model_prefill": float(glat["model_prefill"]),
                "model_decode": float(glat["model_decode"]),
                "end_to_end": float(glat["end_to_end"]),
                "database_retrieve": float(glat["database_retrieval"]),
                "effective_memory_path": float(glat["query_embedding"])
                + float(glat["database_retrieval"]),
                "e2e_provenance": "measured",
            },
            "Polyglot": {
                "query_embedding": float(glat["query_embedding"]),
                "vector_search": float(plat["vector"]),
                "relation_scan": float(plat["seed_relational_filter"])
                + float(plat["graph_relational_filter"]),
                "graph_traversal": float(plat["graph"]),
                "merge_memory": float(plat["merge_and_prompt"]),
                "prompt_processing": float(glat["prompt_serialization"])
                + float(glat["prompt_tokenization"]),
                "model_prefill": float(glat["model_prefill"]),
                "model_decode": float(glat["model_decode"]),
                "end_to_end": max(
                    0.0, float(glat["end_to_end"]) - float(glat["database_retrieval"])
                )
                + float(plat["total"]),
                "database_retrieve": float(plat["total"]),
                "effective_memory_path": float(glat["query_embedding"])
                + float(plat["total"]),
                "e2e_provenance": "reconstructed after exact prompt parity",
            },
        }
        for system, values in systems.items():
            accounted = sum(
                float(values[key])
                for key in (
                    "query_embedding",
                    "vector_search",
                    "relation_scan",
                    "graph_traversal",
                    "merge_memory",
                    "prompt_processing",
                    "model_prefill",
                    "model_decode",
                )
            )
            values["other_orchestration"] = max(
                0.0, float(values["end_to_end"]) - accounted
            )
            per_query.append({"target_id": target, "system": system, **values})

    phase_keys = (
        "query_embedding",
        "vector_search",
        "relation_scan",
        "graph_traversal",
        "merge_memory",
        "prompt_processing",
        "model_prefill",
        "model_decode",
        "other_orchestration",
        "end_to_end",
    )
    summaries: list[dict[str, Any]] = []
    for system in ARM_COLORS:
        rows = [row for row in per_query if row["system"] == system]
        for phase in phase_keys:
            stats = _dist(float(row[phase]) for row in rows)
            summaries.append({"system": system, "phase": phase, **stats})
    return per_query, summaries


def _paper_table5_rows(metrics: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for arm in ("memory_off", "full_gem", "polyglot"):
        source = metrics["quality"][arm]["paper_difficulty_table"]
        for category, tiers in source.items():
            for difficulty, cell in tiers.items():
                rows.append(
                    {
                        "System": ARM_LABELS[arm],
                        "Category": category,
                        "Difficulty": difficulty,
                        "Accuracy (%)": 100.0 * float(cell["accuracy"]),
                        "Contexts": cell.get("contexts"),
                        "Provenance": metrics["quality"][arm]["provenance"],
                    }
                )
    return rows


def _style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for column in sheet.columns:
            width = min(
                70, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
            sheet.column_dimensions[column[0].column_letter].width = width


def _append_rows(
    sheet: Any, headers: Sequence[str], rows: Iterable[Sequence[Any]]
) -> None:
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))


def write_workbook(
    output: Path,
    metrics: dict[str, Any],
    table5_rows: list[dict[str, Any]],
    per_query: list[dict[str, Any]],
    phase_rows: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    _append_rows(
        readme,
        ("Field", "Value"),
        (
            ("Generated at", datetime.now(timezone.utc).isoformat()),
            ("Paper", "EvoMemBench, arXiv:2605.18421"),
            (
                "Paper Table 3 metric",
                "Answer accuracy (%); average rank; lower rank is better",
            ),
            ("Paper Table 5 metric", "Answer accuracy (%) by category and difficulty"),
            ("Local Table 5 status", "COMPLETE: 120 contexts / 884 episodes"),
            (
                "Local Table 3 status",
                "NOT RUN in this workbook; no local values are imputed",
            ),
            (
                "Strict Table 3 reproducibility",
                "NO on current stack: paper models/judge differ and raw contexts exceed local 262K limit",
            ),
            (
                "Controlled Table 3 reproducibility",
                "YES after declaring a 262K context-budget policy and local judge/model provenance",
            ),
            (
                "Polyglot quality",
                "Inherited from GEM only after 100% ordered-ID and injection-SHA parity",
            ),
            (
                "Polyglot E2E latency",
                "Reconstructed from measured GEM generation path plus measured Polyglot retrieval",
            ),
            ("Hardware", str(metrics.get("hardware_claim"))),
        ),
    )

    paper = workbook.create_sheet("Paper_Table3")
    _append_rows(paper, PAPER_TABLE3_HEADERS, PAPER_TABLE3)

    local = workbook.create_sheet("Local_Table3_Target")
    _append_rows(
        local,
        PAPER_TABLE3_HEADERS + ("Status", "Context policy", "Provenance"),
        (
            (
                "Control",
                "No Memory",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "NOT_RUN",
                "262K raw-context budget required",
                "local controlled reproduction",
            ),
            (
                "Tri-modal Memory",
                "GEM",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "NOT_RUN",
                "retrieved-memory budget",
                "must be independently generated",
            ),
            (
                "Polyglot Memory",
                "Polyglot",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "NOT_RUN",
                "retrieved-memory budget",
                "quality may be inherited only after prompt parity",
            ),
        ),
    )

    coverage = workbook.create_sheet("INEP_Coverage")
    _append_rows(
        coverage,
        (
            "Table 3 column",
            "Dimension",
            "Questions",
            "Official subset",
            "Official quality metric",
            "Min raw context tokens",
            "Max raw context tokens",
            "Fits local 262K raw baseline",
        ),
        INEP_COVERAGE,
    )

    table5 = workbook.create_sheet("Local_Table5")
    headers = (
        "System",
        "Category",
        "Difficulty",
        "Accuracy (%)",
        "Contexts",
        "Provenance",
    )
    _append_rows(
        table5, headers, ([row[key] for key in headers] for row in table5_rows)
    )

    latency = workbook.create_sheet("Latency_Summary")
    latency_rows = []
    for system, key in (
        ("No Memory", "no_memory_end_to_end_measured"),
        ("GEM", "gem_end_to_end_measured"),
        ("Polyglot", "polyglot_end_to_end_parity_reconstructed"),
    ):
        value = summary["latency_ms"][key]
        latency_rows.append(
            (
                system,
                key,
                value["n"],
                value["mean"],
                value["p50"],
                value["p95"],
                value["p99"],
            )
        )
    for system in ("GEM", "Polyglot"):
        rows = [row for row in per_query if row["system"] == system]
        for field, label in (
            ("database_retrieve", "database_retrieve_measured"),
            (
                "effective_memory_path",
                "query_embedding_plus_database_retrieve_measured",
            ),
        ):
            value = _dist(float(row[field]) for row in rows)
            latency_rows.append(
                (
                    system,
                    label,
                    value["n"],
                    value["mean"],
                    value["p50"],
                    value["p95"],
                    value["p99"],
                )
            )
    _append_rows(
        latency,
        ("System", "Metric", "N", "Mean ms", "P50 ms", "P95 ms", "P99 ms"),
        latency_rows,
    )

    phases = workbook.create_sheet("Phase_Breakdown")
    phase_headers = ("system", "phase", "n", "mean", "p50", "p95", "p99")
    _append_rows(
        phases,
        phase_headers,
        ([row[key] for key in phase_headers] for row in phase_rows),
    )

    _style_workbook(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def render_latency(
    summary: dict[str, Any],
    per_query: list[dict[str, Any]],
    phase_rows: list[dict[str, Any]],
    output_dir: Path,
) -> None:
    import matplotlib.pyplot as plt

    output_dir.mkdir(parents=True, exist_ok=True)
    systems = list(ARM_COLORS)
    latency_keys = (
        "no_memory_end_to_end_measured",
        "gem_end_to_end_measured",
        "polyglot_end_to_end_parity_reconstructed",
    )
    figure, axes = plt.subplots(1, 2, figsize=(13.2, 5.2), constrained_layout=True)

    p50 = [summary["latency_ms"][key]["p50"] / 1000.0 for key in latency_keys]
    p95 = [summary["latency_ms"][key]["p95"] / 1000.0 for key in latency_keys]
    x = range(len(systems))
    axes[0].bar(
        [value - 0.18 for value in x], p50, width=0.36, label="P50", color="#4C78A8"
    )
    axes[0].bar(
        [value + 0.18 for value in x], p95, width=0.36, label="P95", color="#E45756"
    )
    axes[0].set_xticks(list(x), systems)
    axes[0].set_ylabel("End-to-end latency (s)")
    axes[0].set_title("Three-system latency")
    axes[0].legend(frameon=False)

    phases = (
        "query_embedding",
        "vector_search",
        "relation_scan",
        "graph_traversal",
        "merge_memory",
        "prompt_processing",
        "model_prefill",
        "model_decode",
        "other_orchestration",
    )
    phase_labels = {
        "query_embedding": "Query embedding",
        "vector_search": "Vector search",
        "relation_scan": "Relation scan",
        "graph_traversal": "Graph traversal",
        "merge_memory": "Fused retrieval / merge",
        "prompt_processing": "Prompt processing",
        "model_prefill": "Model prefill",
        "model_decode": "Model decode",
        "other_orchestration": "Other orchestration",
    }
    phase_colors = (
        "#72B7B2",
        "#54A24B",
        "#B279A2",
        "#9D755D",
        "#FF9DA6",
        "#BAB0AC",
        "#F2CF5B",
        "#E45756",
        "#A0CBE8",
    )
    lookup = {
        (row["system"], row["phase"]): float(row["mean"]) / 1000.0 for row in phase_rows
    }
    bottoms = [0.0] * len(systems)
    for phase, color in zip(phases, phase_colors, strict=True):
        values = [lookup[(system, phase)] for system in systems]
        axes[1].bar(
            systems, values, bottom=bottoms, label=phase_labels[phase], color=color
        )
        bottoms = [
            bottom + value for bottom, value in zip(bottoms, values, strict=True)
        ]
    axes[1].set_ylabel("Mean end-to-end latency (s)")
    axes[1].set_title("Mean phase-time breakdown")
    axes[1].legend(
        frameon=False, fontsize=8, bbox_to_anchor=(1.02, 1), loc="upper left"
    )

    figure.suptitle(
        "EvoMemBench CrossEp-Know latency (764 reuse decisions)",
        fontsize=14,
        fontweight="bold",
    )
    for suffix in ("png", "pdf", "svg"):
        figure.savefig(
            output_dir / f"three_system_latency_and_phase_breakdown.{suffix}", dpi=300
        )
    plt.close(figure)

    retrieval_figure, retrieval_axis = plt.subplots(
        figsize=(8.8, 5.4), constrained_layout=True
    )
    retrieval_phases = (
        "query_embedding",
        "vector_search",
        "relation_scan",
        "graph_traversal",
        "merge_memory",
    )
    retrieval_colors = ("#72B7B2", "#54A24B", "#B279A2", "#9D755D", "#FF9DA6")
    retrieval_bottoms = [0.0] * len(systems)
    for phase, color in zip(retrieval_phases, retrieval_colors, strict=True):
        values = [lookup[(system, phase)] * 1000.0 for system in systems]
        retrieval_axis.bar(
            systems,
            values,
            bottom=retrieval_bottoms,
            label=phase_labels[phase],
            color=color,
        )
        retrieval_bottoms = [
            bottom + value
            for bottom, value in zip(retrieval_bottoms, values, strict=True)
        ]
    retrieval_axis.set_ylabel("Mean memory-path latency (ms)")
    retrieval_axis.set_title(
        "Memory-path phase breakdown\n"
        "GEM database stages are fused; Polyglot stages are separately measured"
    )
    retrieval_axis.legend(frameon=False, fontsize=9)
    retrieval_axis.grid(axis="y", alpha=0.2)
    for suffix in ("png", "pdf", "svg"):
        retrieval_figure.savefig(
            output_dir / f"three_system_memory_path_breakdown.{suffix}", dpi=300
        )
    plt.close(retrieval_figure)

    retrieve_figure, retrieve_axes = plt.subplots(
        1, 2, figsize=(11.8, 5.0), constrained_layout=True
    )
    memory_systems = ("GEM", "Polyglot")
    for axis, field, title in (
        (
            retrieve_axes[0],
            "database_retrieve",
            "Database retrieve latency\n(query vector already available)",
        ),
        (
            retrieve_axes[1],
            "effective_memory_path",
            "Effective agent-memory path\n(query embedding + retrieve)",
        ),
    ):
        distributions = []
        for system in memory_systems:
            rows = [row for row in per_query if row["system"] == system]
            distributions.append(_dist(float(row[field]) for row in rows))
        positions = range(len(memory_systems))
        axis.bar(
            [value - 0.18 for value in positions],
            [row["p50"] for row in distributions],
            width=0.36,
            label="P50",
            color="#4C78A8",
        )
        axis.bar(
            [value + 0.18 for value in positions],
            [row["p95"] for row in distributions],
            width=0.36,
            label="P95",
            color="#E45756",
        )
        axis.set_xticks(list(positions), memory_systems)
        axis.set_ylabel("Latency (ms)")
        axis.set_title(title)
        axis.grid(axis="y", alpha=0.2)
        axis.legend(frameon=False)
    retrieve_figure.suptitle(
        "Agent-memory retrieval latency (764 paired reuse decisions)",
        fontsize=14,
        fontweight="bold",
    )
    for suffix in ("png", "pdf", "svg"):
        retrieve_figure.savefig(
            output_dir / f"three_system_retrieve_latency.{suffix}", dpi=300
        )
    plt.close(retrieve_figure)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-root", type=Path, required=True)
    parser.add_argument("--poly-root", type=Path, required=True)
    parser.add_argument("--metrics", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--figure-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    metrics = _json(args.metrics)
    summary = _json(args.summary)
    if (
        summary.get("status") != "complete"
        or metrics.get("parity", {}).get("fraction") != 1.0
    ):
        raise ValueError("the three-system source run is incomplete or failed parity")
    per_query, phase_rows = _phase_rows(args.run_root, args.poly_root)
    table5_rows = _paper_table5_rows(metrics)
    write_workbook(args.xlsx, metrics, table5_rows, per_query, phase_rows, summary)
    render_latency(summary, per_query, phase_rows, args.figure_dir)
    receipt = {
        "schema_version": "evomembench_table3_table5_export_v0.1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "xlsx": str(args.xlsx.resolve()),
        "figure_dir": str(args.figure_dir.resolve()),
        "local_table5_rows": len(table5_rows),
        "phase_summary_rows": len(phase_rows),
        "table3_paper_rows": len(PAPER_TABLE3),
        "local_table3_status": "not_run",
    }
    (args.figure_dir / "export_receipt.json").write_text(
        json.dumps(receipt, indent=2) + "\n"
    )


if __name__ == "__main__":
    main()
